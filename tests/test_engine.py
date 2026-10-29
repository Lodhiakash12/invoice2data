"""Engine tests. Run from the project folder:  python -m pytest -q tests

- every template passes the load-time checks
- every sample PDF still gives its expected output (tests/expected/*.json)
- a new field (cess, round-off) works with a YAML-only change
- template mistakes are rejected at load (including bad replacement strings)
- template versions: the version with the most evidence represents the layout
- damaged values are caught by the checks; checks on all-zero values prove nothing
- rules cannot run code
- look-alike templates: chosen by evidence, never by name; ambiguous cases go to review
- audit defects: slow patterns are stopped, one bad file does not stop the run, copies and
  unreadable numbers do not merge documents, outputs do not overwrite each other, scanned pages
  are flagged
- document types: a new type is only a folder; routing keeps types apart; the keyword index
  finds the same templates as a full scan; the overlap test catches competing templates
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

import pytest
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from invoice_engine import Registry, layout, load_catalog, load_templates, poppler, process_pdf  # noqa: E402
from invoice_engine.document import process_group  # noqa: E402
from invoice_engine.rules import COMMON_FUNCTIONS, RuleError, parse_rule  # noqa: E402
from invoice_engine.segment import matching_templates, page_texts  # noqa: E402

CATALOG = load_catalog(ROOT)
INVOICE = CATALOG.type_named("invoice")
INVOICE_TEMPLATES = ROOT / "document_types" / "invoice" / "templates"
REGISTRY = INVOICE.registry                      # invoice fields: shared + document_types/invoice
TEMPLATES = INVOICE.template_set
EXPECTED = sorted((ROOT / "tests" / "expected").glob("*.json"))


def run_text(text: str, templates: list, pdf: Path | None = None) -> dict:
    """Parse one page of text the way process_pdf does (pdf: the file the text came from, for
    the fields read from word positions)."""
    candidates = matching_templates(text, templates)
    assert candidates, "no template matched"
    group = {"candidates": candidates, "texts": [text], "pages": [1]}
    return process_group(group, pdf or ROOT / "test.pdf", REGISTRY)


def write_template(folder: Path, name: str, body: str) -> None:
    (folder / name).write_text(body, encoding="utf-8")


# ----------------------------------------------------------------------------- templates

def test_all_templates_load_without_errors():
    assert CATALOG.rejected == {}, CATALOG.rejected
    assert len(TEMPLATES.templates) >= 14 and CATALOG.type_named("tax_document").template_set.templates


@pytest.mark.parametrize("expected_file", EXPECTED, ids=[p.stem for p in EXPECTED])
def test_sample_pdfs_match_expected_output(expected_file: Path):
    expected = json.loads(expected_file.read_text(encoding="utf-8"))
    pdf = ROOT / expected["source_file"]
    if not pdf.exists():
        pytest.skip(f"{pdf.name} not in the project folder")
    result = json.loads(json.dumps(process_pdf(pdf, CATALOG), default=str))
    assert len(result["documents"]) == len(expected["documents"])
    for got, want in zip(result["documents"], expected["documents"]):
        for key in ("pages", "template", "status", "fields", "lines"):
            assert got[key] == want[key], f"{pdf.name} pages {want['pages']}: {key} changed"


# ------------------------------------------------------------- new fields need YAML only

CESS_TEMPLATE = r"""
issuer: Acme Foods
keywords: ['ACME FOODS PVT LTD', 'TAX INVOICE']
required_fields: [issuer, document_type, invoice_number, date, amount, seller_name]
fields:
  document_type: {parser: static, value: tax_invoice}
  invoice_number: 'Invoice No:\s*(\S+)'
  date: 'Invoice Date:\s*(\d{2}/\d{2}/\d{4})'
  seller_name: '\A\s*(ACME FOODS PVT LTD)'
  amount_untaxed: 'Taxable value\s+([\d.]+)'
  amount_cgst: '(?m)^CGST\s+([\d.]+)'
  amount_sgst: '(?m)^SGST\s+([\d.]+)'
  amount_cess: '(?m)^Cess\s+([\d.]+)'
  amount_round_off: 'Round off\s+(-?[\d.]+)'
  amount: 'Grand total\s+([\d.]+)'
  lines:
    parser: lines
    start: 'Item\s+Qty'
    end: 'Taxable value'
    first_line: '^(?P<name>\S.*?)\s{2,}(?P<qty>\d+)\s+(?P<price_unit>[\d.]+)\s+(?P<price_subtotal>[\d.]+)\s+(?P<cgst_rate>[\d.]+)%\s+(?P<cgst_amount>[\d.]+)\s+(?P<sgst_rate>[\d.]+)%\s+(?P<sgst_amount>[\d.]+)\s+(?P<cess_rate>[\d.]+)%\s+(?P<cess_amount>[\d.]+)\s+(?P<price_total>[\d.]+)\s*$'
    line: '^\s+(?P<name>\S.*?)\s*$'
    types: {qty: int, price_unit: float, price_subtotal: float, cgst_rate: float, cgst_amount: float,
            sgst_rate: float, sgst_amount: float, cess_rate: float, cess_amount: float, price_total: float}
options: {currency: INR, date_formats: ['%d/%m/%Y']}
"""

CESS_TEXT = """ACME FOODS PVT LTD
TAX INVOICE
Invoice No: AF-001
Invoice Date: 05/10/2026

Item              Qty    Rate   Taxable   CGST          SGST          Cess          Total
Masala Tea          3   66.50    199.50   2.5%   4.99   2.5%   4.99   1.0%   2.00   211.48

Taxable value   199.50
CGST            4.99
SGST            4.99
Cess            2.00
Round off       0.52
Grand total     212.00
"""


def test_new_tax_and_round_off_work_without_code_changes(tmp_path: Path):
    write_template(tmp_path, "acme.yml", CESS_TEMPLATE)
    templates = load_templates(tmp_path, REGISTRY)
    assert templates.rejected == {}
    doc = run_text(CESS_TEXT, templates.templates)
    assert doc["status"] == "OK", doc["issues"]
    assert doc["fields"]["amount_cess"] == 2.0 and doc["fields"]["amount_round_off"] == 0.52


def test_wrong_cess_amount_is_caught(tmp_path: Path):
    write_template(tmp_path, "acme.yml", CESS_TEMPLATE)
    templates = load_templates(tmp_path, REGISTRY)
    doc = run_text(CESS_TEXT.replace("Cess            2.00", "Cess            3.00"), templates.templates)
    assert doc["status"] == "FAIL"


# ------------------------------------------------------------------- load-time checks

def test_template_mistakes_are_rejected(tmp_path: Path):
    write_template(tmp_path, "bad.yml", r"""
issuer: Bad
keywords: [Bad Ltd, BILL TO: Someone]
required_fields: [issuer, invoice_number, date, amount, seller_name]
fields:
  invoice_number: 'Invoice (No|Number):\s*(\S+)'
  date: 'Date:\s*(\d{2}-\d{2}-\d{4})'
  amount_tcs: 'TCS\s+([\d.]+)'
  amount: '^Total\s+([\d.]+)'
lines:
  start: 'Item'
  end: '^Total'
  line: '^(?P<name>.+)$'
checks:
  document: [{rule: "__import__('os') == 1"}]
""")
    templates = load_templates(tmp_path, REGISTRY)
    problems = " ".join(templates.rejected["bad.yml"])
    assert "not text" in problems                       # unquoted 'BILL TO:' became a mapping
    assert "exactly 1 capture group" in problems
    assert "amount field is not in config/fields.yml" in problems
    assert "without (?m)" in problems
    assert "date_formats" in problems
    assert "seller_name" in problems                    # required but never defined
    assert "not allowed" in problems                    # code in a rule


def test_rules_cannot_run_code():
    for text in ("__import__('os').system('dir') == 0", "amount.real == 1", "[1][0] == 1", "lambda: 1 == 1"):
        with pytest.raises(RuleError):
            parse_rule(text, COMMON_FUNCTIONS)


# ------------------------------------------------------------------------- versions

def test_newer_version_is_tried_first_and_older_is_the_fallback(tmp_path: Path):
    old = (INVOICE_TEMPLATES / "anthropic_invoice.yml").read_text(encoding="utf-8")
    new = ("family: anthropic_invoice\nversion: 2\n"
           + old.replace(r"amount: '(?m)^\s+Total[ \t]{2,}", r"amount: '(?m)^\s+Total amount[ \t]{2,}"))
    write_template(tmp_path, "anthropic_invoice.yml", old)
    write_template(tmp_path, "anthropic_invoice_v2.yml", new)
    templates = load_templates(tmp_path, REGISTRY)
    assert templates.rejected == {}
    pdf = ROOT / "Invoice-EXQCNVGN-0004.pdf"
    if not pdf.exists():
        pytest.skip("Anthropic sample PDF not in the project folder")
    old_layout = page_texts(pdf)[0]
    new_layout = old_layout.replace("    Total   ", "    Total amount", 1)

    doc = run_text(old_layout, templates.templates, pdf)     # old PDF: v2 fails, v1 used
    assert (doc["version"], doc["status"]) == (1, "OK")
    assert any("also tried anthropic_invoice_v2.yml" in i["message"] for i in doc["issues"])

    doc = run_text(new_layout, templates.templates, pdf)     # new PDF: v2 used directly
    assert (doc["version"], doc["status"]) == (2, "OK")


def test_duplicate_version_is_rejected(tmp_path: Path):
    body = (INVOICE_TEMPLATES / "anthropic_invoice.yml").read_text(encoding="utf-8")
    write_template(tmp_path, "a.yml", "family: anthropic_invoice\n" + body)
    write_template(tmp_path, "b.yml", "family: anthropic_invoice\n" + body)
    templates = load_templates(tmp_path, REGISTRY)
    assert "b.yml" in templates.rejected


# ------------------------------------------------------- look-alike templates (regression)
# Each test reproduces one failure found in the look-alike investigation (rnd/eval in the scratch
# folder): before the change the engine chose by alphabetical order or accepted unverified totals.

def derive(folder: Path, source: str, target: str, **changes) -> None:
    """Write templates/<source> with some keys changed (None removes a key) as <folder>/<target>."""
    data = yaml.safe_load((INVOICE_TEMPLATES / source).read_text(encoding="utf-8"))
    for key, value in changes.items():
        if value is None:
            data.pop(key, None)
        else:
            data[key] = value
    (folder / target).write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def pool_with(folder: Path, *extra: str) -> list:
    """All real invoice templates plus the extra files already written to folder."""
    for path in (INVOICE_TEMPLATES).glob("*.yml"):
        if not (folder / path.name).exists():
            shutil.copy(path, folder / path.name)
    templates = load_templates(folder, REGISTRY)
    assert templates.rejected == {}, templates.rejected
    return templates.templates


def need(pdf_name: str) -> Path:
    pdf = ROOT / pdf_name
    if not pdf.exists():
        pytest.skip(f"{pdf_name} not in the project folder")
    return pdf


def test_identical_lookalike_is_sent_to_review_not_chosen_by_name(tmp_path: Path):
    pdf = need("ms_invoice_326255343871054720_1.pdf")
    derive(tmp_path, "meesho_tax_invoice.yml", "abc_marketplace_invoice.yml", issuer="ABC Marketplace")
    doc = process_pdf(pdf, pool_with(tmp_path), REGISTRY)["documents"][0]
    assert doc["status"] == "CHECK"
    assert any(i["level"] == "WARN" and i["message"].startswith("ambiguous:") for i in doc["issues"])
    assert [a["template"] for a in doc["alternatives"]] == ["meesho_tax_invoice.yml"]


def test_wrong_template_that_does_not_fail_loses_to_better_evidence(tmp_path: Path):
    pdf = need("OD437334801620974100.pdf")
    shared = ["Tax Invoice", "Sold By:", "Invoice Number #"]
    fee = yaml.safe_load((INVOICE_TEMPLATES / "flipkart_platform_fee_invoice.yml").read_text(encoding="utf-8"))
    fee["fields"].pop("lines")                       # a look-alike without an item table never FAILs on it
    derive(tmp_path, "flipkart_seller_tax_invoice.yml", "flipkart_seller_tax_invoice.yml",
           keywords=shared, exclude_keywords=None)
    derive(tmp_path, "flipkart_platform_fee_invoice.yml", "flipkart_platform_fee_invoice.yml",
           keywords=shared, exclude_keywords=None, fields=fee["fields"])
    docs = process_pdf(pdf, pool_with(tmp_path), REGISTRY)["documents"]
    by_pages = {tuple(d["pages"]): d for d in docs}
    for pages in ((1,), (8,)):
        assert (by_pages[pages]["template"], by_pages[pages]["status"]) == ("flipkart_seller_tax_invoice.yml", "OK")


def test_evidence_poor_duplicate_does_not_win_by_alphabetical_order(tmp_path: Path):
    pdf = need("oyo.pdf")
    (tmp_path / "com.oyo.invoice.yml").write_text(          # like invoice2data's own OYO template
        "issuer: OYO\nkeywords: [OYO, Oravel, Stays]\n"
        "fields:\n"
        "  amount: 'Grand Total\\s+Rs (\\d+)'\n"
        "  date: 'Date:\\s(\\d{1,2}/\\d{1,2}/\\d{4})'\n"
        "  invoice_number: '([A-Z0-9]+)\\s+Cash at Hotel'\n"
        "options: {currency: INR, date_formats: ['%d/%m/%Y']}\n", encoding="utf-8")
    doc = process_pdf(pdf, pool_with(tmp_path), REGISTRY)["documents"][0]
    assert (doc["template"], doc["status"]) == ("oyo_payment_receipt.yml", "OK")
    assert any("also tried com.oyo.invoice.yml" in i["message"] for i in doc["issues"])


GENERIC = r"""
issuer: Generic
keywords: [Invoice]
required_fields: [issuer, invoice_number, date, amount]
fields:
  invoice_number: '(?i)invoice\s*(?:number|no\.?)\s*[:#]?\s*([A-Z0-9][A-Z0-9-]{3,})'
  date: '(?i)(?:date of issue|invoice date)\s*:?\s*([A-Z][a-z]+ \d{1,2}, \d{4})'
  amount: '(?m)^\s*(?:Grand\s+)?Total\s+\D{0,4}([\d,]+\.\d{2})\s*$'
options: {currency: INR}
"""


def test_total_that_no_check_verified_is_not_accepted(tmp_path: Path):
    pdf = need("Invoice-EXQCNVGN-0004.pdf")
    write_template(tmp_path, "generic_invoice.yml", GENERIC)
    templates = load_templates(tmp_path, REGISTRY)          # the Anthropic template is not there
    doc = process_pdf(pdf, templates.templates, REGISTRY)["documents"][0]
    assert doc["fields"]["amount"] == 2399.0 and doc["evidence"]["checks_passed"] == 0
    assert doc["status"] == "CHECK"
    assert any("amount is not verified" in i["message"] for i in doc["issues"])


def test_catch_all_template_never_beats_the_real_one(tmp_path: Path):
    write_template(tmp_path, "generic_invoice.yml", GENERIC)
    templates = pool_with(tmp_path)
    for name in ("Invoice-EXQCNVGN-0004.pdf", "Invoice_8626225731.pdf", "ms_invoice_326255343871054720_1.pdf"):
        pdf = need(name)
        for doc in process_pdf(pdf, templates, REGISTRY)["documents"]:
            assert doc["template"] != "generic_invoice.yml" and doc["status"] == "OK", (name, doc["pages"])


def test_lookalike_matching_page_one_only_does_not_split_the_invoice(tmp_path: Path):
    pdf = need("QualityHosting.pdf")
    base = yaml.safe_load((INVOICE_TEMPLATES / "qualityhosting_custom.yml").read_text(encoding="utf-8"))
    base["fields"].pop("lines")
    derive(tmp_path, "qualityhosting_custom.yml", "aaa_hosting_invoice.yml", issuer="AAA Hosting",
           keywords=["QualityHosting AG", "Contract No."], fields=base["fields"])
    result = process_pdf(pdf, pool_with(tmp_path), REGISTRY)
    assert [(d["pages"], d["template"], d["status"]) for d in result["documents"]] == \
        [([1, 2], "qualityhosting_custom.yml", "OK")]


def test_lookalike_on_all_pages_does_not_hide_the_real_template(tmp_path: Path):
    pdf = need("OD437334801620974100.pdf")       # bill of supply on pages 2-3; page 3 matches no real template
    base = yaml.safe_load((INVOICE_TEMPLATES / "flipkart_bill_of_supply.yml").read_text(encoding="utf-8"))
    base["fields"].pop("lines")
    derive(tmp_path, "flipkart_bill_of_supply.yml", "aaa_supply_note.yml", issuer="AAA Supply",
           keywords=["VILLAGE SANPKA", "DISTRICT GURUGRAM"], fields=base["fields"], continuation_keywords=None)
    docs = process_pdf(pdf, pool_with(tmp_path), REGISTRY)["documents"]
    doc = next(d for d in docs if d["pages"] == [2, 3])
    assert (doc["template"], doc["status"]) == ("flipkart_bill_of_supply.yml", "OK")


def test_shared_layouts_do_not_chain_separate_invoices(tmp_path: Path):
    pdf = need("Invoice_8626225731.pdf")
    for name in ("zomato_restaurant_invoice.yml", "zomato_platform_fee_invoice.yml"):
        derive(tmp_path, name, name, keywords=["Invoice No"], exclude_keywords=None)
    result = process_pdf(pdf, pool_with(tmp_path), REGISTRY)
    assert [(d["pages"], d["template"], d["status"]) for d in result["documents"]] == \
        [([1], "zomato_restaurant_invoice.yml", "OK"), ([2], "zomato_platform_fee_invoice.yml", "OK")]


# ---------------------------------------------------------- layout fingerprint (tie-breaker)

def test_fingerprint_keeps_label_words_only_and_ignores_rows():
    page = [("Invoice", 0.62, 100.0), ("Nexora", 0.70, 90.0), ("Subtotal", 0.61, 500.0), ("18,500.00", 0.9, 500.0)]
    order = layout.fingerprint([page], {"invoice", "subtotal"})
    assert order == [["invoice", 7], ["subtotal", 7]]                   # no names, no amounts
    longer_table = [(w, x, y + 300 if y > 200 else y) for w, x, y in page]   # totals pushed down
    assert layout.similarity(layout.fingerprint([longer_table], {"invoice", "subtotal"}), order) == 1.0
    moved = [(w, 0.05, y) for w, x, y in page]                           # same words, other column
    assert layout.similarity(layout.fingerprint([moved], {"invoice", "subtotal"}), order) == 0.0


def test_malformed_layout_block_is_rejected(tmp_path: Path):
    derive(tmp_path, "oyo_payment_receipt.yml", "oyo_payment_receipt.yml",
           layout={"columns": 12, "sample_pages": 1, "order": [["receipt", 40]]})
    assert "oyo_payment_receipt.yml" in load_templates(tmp_path, REGISTRY).rejected


def _store_layout(folder: Path, name: str, pdf: Path, pages: list[int]) -> None:
    templates = {t["template_name"]: t for t in load_templates(folder, REGISTRY).templates}
    order = layout.fingerprint(layout.page_words(pdf, pages), layout.label_vocabulary(templates[name]))
    layout.write_layout(folder / name, layout.layout_block(order, 1))


def test_layout_breaks_a_tie_only_when_one_layout_is_clearly_closer(tmp_path: Path):
    meesho, flipkart = need("ms_invoice_326255343871054720_1.pdf"), need("OD437334801620974100.pdf")
    derive(tmp_path, "meesho_tax_invoice.yml", "abc_marketplace_invoice.yml", issuer="ABC Marketplace")
    pool_with(tmp_path)
    _store_layout(tmp_path, "meesho_tax_invoice.yml", meesho, [1])
    _store_layout(tmp_path, "abc_marketplace_invoice.yml", flipkart, [1])    # ABC's sample looks different
    doc = process_pdf(meesho, load_templates(tmp_path, REGISTRY).templates, REGISTRY)["documents"][0]
    assert (doc["template"], doc["status"]) == ("meesho_tax_invoice.yml", "OK")
    assert any(i["message"].startswith("layout decided the tie") for i in doc["issues"])

    _store_layout(tmp_path, "abc_marketplace_invoice.yml", meesho, [1])      # same layout: no decision
    doc = process_pdf(meesho, load_templates(tmp_path, REGISTRY).templates, REGISTRY)["documents"][0]
    assert doc["status"] == "CHECK"
    assert any(i["message"].startswith("ambiguous:") for i in doc["issues"])


# --------------------------------------------------------------- finding and running pdftotext

def _blank_pdf(path: Path) -> None:
    """A valid one-page PDF with no text at all (like a scan without OCR)."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 4\n0000000000 65535 f \n" + b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref
    path.write_bytes(bytes(out))


def test_pdftotext_is_found_from_settings_without_the_path(tmp_path: Path, monkeypatch):
    real = poppler.configure(None)                           # wherever it is on this machine
    (tmp_path / "settings.yml").write_text(f"poppler_path: {Path(real.path).parent}\n", encoding="utf-8")
    monkeypatch.setenv("PATH", "")
    monkeypatch.delenv("POPPLER_PATH", raising=False)
    found = poppler.configure(tmp_path)
    assert (found.source, found.version) == ("settings.yml", real.version)
    monkeypatch.setenv("POPPLER_PATH", str(Path(real.path).parent))   # the environment variable wins
    assert poppler.configure(tmp_path).source == "POPPLER_PATH"


def test_missing_pdftotext_explains_where_it_looked(tmp_path: Path, monkeypatch):
    (tmp_path / "settings.yml").write_text("poppler_path: C:\\nowhere\\bin\n", encoding="utf-8")
    monkeypatch.setenv("PATH", "")
    monkeypatch.delenv("POPPLER_PATH", raising=False)
    with pytest.raises(poppler.PopplerError) as exc:
        poppler.configure(tmp_path)
    assert "nowhere" in str(exc.value) and "settings.yml" in str(exc.value)


def test_unreadable_pdfs_are_reported_not_crashed(tmp_path: Path):
    poppler.configure(None)
    damaged = tmp_path / "damaged.pdf"
    damaged.write_bytes(b"this is not a pdf at all")
    blank = tmp_path / "scan.pdf"
    _blank_pdf(blank)
    result = process_pdf(damaged, TEMPLATES.templates, REGISTRY)
    assert result["documents"] == [] and "damaged or is not a real PDF" in result["notes"][0]
    result = process_pdf(blank, TEMPLATES.templates, REGISTRY)
    assert any("no page has a text layer" in n for n in result["notes"])
    assert "password-protected" in poppler._explain("Command Line Error: Incorrect password")


# ------------------------------------------------------------------- damaged values

DAMAGE = [
    ("Invoice_8626225731.pdf", 0, "10.38", "18.38"),
    ("Invoice_8626225731.pdf", 0, "99.89", "98.89"),
    ("Invoice_8626225731.pdf", 1, "Platform fee", "Platform\nfee"),
    ("Invoice_8626225731.pdf", 1, "Invoice No:", "Invoice Number:"),
    ("ms_invoice_326255343871054720_1.pdf", 0, ":Rs.8.70", ":Rs.8.07"),
    ("Invoice-EXQCNVGN-0004.pdf", 0, "365.95", "356.95"),
    ("OD437334801620974100.pdf", 0, "30.95", "39.95"),
]


@pytest.mark.parametrize("pdf_name,page,old,new", DAMAGE)
def test_damaged_values_are_caught(pdf_name, page, old, new):
    pdf = ROOT / pdf_name
    if not pdf.exists():
        pytest.skip(f"{pdf_name} not in the project folder")
    text = page_texts(pdf)[page]
    assert old in text
    doc = run_text(text.replace(old, new, 1), TEMPLATES.templates)
    assert doc["status"] == "FAIL", f"not caught: {old} -> {new}"


# ------------------------------------------------- audit defects (each reproduced, then fixed)

from invoice_engine import safe_regex  # noqa: E402
from invoice_engine.export import output_names, write_json  # noqa: E402
from invoice_engine.segment import page_mark, split_into_documents  # noqa: E402


def test_slow_pattern_is_stopped_and_only_its_template_fails(tmp_path: Path):
    """D6: a pattern with exponential run time used to hang the whole run."""
    write_template(tmp_path, "acme.yml", CESS_TEMPLATE)
    slow = CESS_TEMPLATE.replace("issuer: Acme Foods", "issuer: Acme Slow").replace(
        "  amount_cess:", "  customer_number: '(?m)^Ref:\\s*((?:a|aa)+)$'\n  amount_cess:")
    write_template(tmp_path, "acme_slow.yml", slow)
    templates = load_templates(tmp_path, REGISTRY).templates
    old = safe_regex.timeout()
    safe_regex.install(0.5)
    try:
        import time
        started = time.time()
        doc = run_text(CESS_TEXT + "Ref: " + "a" * 60 + "b\n", templates)
        assert time.time() - started < 20
    finally:
        safe_regex.install(old)
    assert (doc["template"], doc["status"]) == ("acme.yml", "OK")
    assert any("acme_slow.yml" in i["message"] and "FAIL" in i["message"] for i in doc["issues"])


def test_bad_replacement_is_rejected_at_load(tmp_path: Path):
    """D4: an undefined group in a replacement passed lint and crashed the run on the first match."""
    body = CESS_TEMPLATE.replace("options: {currency: INR,", "options: {replace: [['(a)', '\\\\2']], currency: INR,")
    write_template(tmp_path, "acme.yml", body)
    rejected = load_templates(tmp_path, REGISTRY).rejected
    assert "acme.yml" in rejected and any("replacement" in m for m in rejected["acme.yml"])


def test_one_file_that_breaks_does_not_stop_the_others(tmp_path: Path, monkeypatch):
    """D4: process_invoice.py had no per-file error handling."""
    pdf = need("oyo.pdf")
    for name in ("a.pdf", "b.pdf"):
        shutil.copy(pdf, tmp_path / name)
    import process_invoice
    from invoice_engine.pipeline import Pipeline
    real = Pipeline.process

    def broken(self, path):
        if Path(path).name == "a.pdf":
            raise ValueError("simulated failure")
        return real(self, path)

    monkeypatch.setattr(Pipeline, "process", broken)
    monkeypatch.setattr(process_invoice, "sample_folder", lambda settings: tmp_path / "samples")
    monkeypatch.setattr(sys, "argv", ["process_invoice.py", str(tmp_path), "--output", str(tmp_path / "out")])
    assert process_invoice.main() == 1
    a = json.loads((tmp_path / "out" / "a.json").read_text(encoding="utf-8"))
    b = json.loads((tmp_path / "out" / "b.json").read_text(encoding="utf-8"))
    assert a["notes"][0].startswith("ERROR: processing stopped") and b["documents"][0]["status"] == "OK"


def test_copies_of_one_document_are_reported_separately(tmp_path: Path):
    """D1: original + duplicate copy in one PDF merged into one document that FAILed."""
    pdf = need("oyo.pdf")
    unite = poppler.tool_path("pdfunite")
    if not unite:
        pytest.skip("pdfunite (Poppler) not available")
    import subprocess
    subprocess.run([unite, str(pdf), str(pdf), str(tmp_path / "twice.pdf")], check=True)
    result = process_pdf(tmp_path / "twice.pdf", TEMPLATES.templates, REGISTRY)
    assert [(d["pages"], d["status"]) for d in result["documents"]] == [([1], "OK"), ([2], "OK")]
    assert result["documents"][1]["copy_of"] == [1]
    single = process_pdf(pdf, TEMPLATES.templates, REGISTRY)["totals_by_document_type"]
    assert result["totals_by_document_type"] == single          # the copy is not counted twice


def test_unreadable_number_does_not_join_another_document():
    """D2: a page whose number could not be read was joined to the previous document."""
    page = page_texts(need("Invoice_8626225731.pdf"))[0]
    number = re.search(r"Invoice No\.\s*:\s*(\S+)", page).group(1)
    other = page.replace(number, "ZZ999999")
    damaged = other.replace("Invoice No.", "Invoice N0.")
    assert [g["pages"] for g in split_into_documents([page, other], TEMPLATES.templates)[0]] == [[1], [2]]
    assert [g["pages"] for g in split_into_documents([page, damaged], TEMPLATES.templates)[0]] == [[1], [2]]


def test_page_marks_are_read():
    assert page_mark("....  Page 2 of 5  ...") == (2, 5)
    assert page_mark("Page 1/2") == (1, 2)
    assert page_mark("Seite 3 von 3") == (3, 3)
    assert page_mark("Page 7 of 5") is None and page_mark("no mark here") is None


def test_outputs_of_files_with_the_same_name_do_not_overwrite(tmp_path: Path):
    """D5: a/x.pdf and b/x.pdf both wrote x.json."""
    names = output_names([Path("a/x.pdf"), Path("b/x.pdf"), Path("c/X.pdf")])
    assert sorted(names.values()) == ["X_3", "x", "x_2"]
    for pdf, name in names.items():
        write_json({"source_file": pdf.name, "documents": [], "unrecognized_pages": [], "notes": []}, tmp_path, name)
    assert len(list(tmp_path.glob("*.json"))) == 3


def test_best_version_wins_not_the_newest(tmp_path: Path):
    """D3: a newer version with a warning beat an older version that reads the document cleanly."""
    pdf = need("oyo.pdf")
    data = yaml.safe_load((INVOICE_TEMPLATES / "oyo_payment_receipt.yml").read_text(encoding="utf-8"))
    data.update(family="oyo_payment_receipt", version=2)
    data["fields"]["customer_number"] = "Never printed label:\\s*(\\S+)"
    (tmp_path / "oyo_payment_receipt_v2.yml").write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
    doc = process_pdf(pdf, pool_with(tmp_path), REGISTRY)["documents"][0]
    assert (doc["template"], doc["status"]) == ("oyo_payment_receipt.yml", "OK")


def test_amounts_without_a_passed_check_are_not_ok(tmp_path: Path):
    """D9: a document with amounts could be OK although no check passed."""
    write_template(tmp_path, "acme.yml", CESS_TEMPLATE.replace(
        "  amount_untaxed: 'Taxable value\\s+([\\d.]+)'\n", "").replace(
        "  amount: 'Grand total\\s+([\\d.]+)'\n", "").replace("amount, seller_name]", "seller_name]"))
    body = (tmp_path / "acme.yml").read_text(encoding="utf-8")
    start = body.index("  lines:")
    end = body.index("options:")
    (tmp_path / "acme.yml").write_text(body[:start] + body[end:], encoding="utf-8")
    doc = run_text(CESS_TEXT, load_templates(tmp_path, REGISTRY).templates)
    assert doc["status"] == "CHECK" and any("amounts are not verified" in i["message"] for i in doc["issues"])


def test_checks_on_all_zero_values_are_not_evidence():
    """A check such as 0 + 0 == 0 agrees but proves nothing about the reading."""
    pdf = need("ACK742817410300726.pdf")
    doc = process_pdf(pdf, CATALOG)["documents"][0]
    assert doc["evidence"]["checks_passed"] == 0 and doc["evidence"]["checks_trivial"] == 4
    assert doc["status"] == "CHECK"


def _scan_like_pdf(path: Path, words: str) -> None:
    """One A4 page covered by a 60 x 85 grey picture, with a line of text on top (like a scan
    that a copier stamped with a header)."""
    import zlib
    pixels = zlib.compress(bytes([200]) * (60 * 85))
    content = (b"q 595 0 0 842 0 0 cm /Im1 Do Q BT /F1 10 Tf 40 820 Td (" + words.encode() + b") Tj ET")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
               b"/Resources << /XObject << /Im1 5 0 R >> /Font << /F1 6 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
               b"<< /Type /XObject /Subtype /Image /Width 60 /Height 85 /ColorSpace /DeviceGray "
               b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n%s\nendstream" % (len(pixels), pixels),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 7\n0000000000 65535 f \n" + b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size 7 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref
    path.write_bytes(bytes(out))


def test_scanned_page_with_a_little_text_is_flagged(tmp_path: Path):
    """D8: a scanned page with a stamped header was only reported as unrecognized."""
    poppler.configure(None)
    if not poppler.tool_path("pdfimages"):
        pytest.skip("pdfimages (Poppler) not available")
    _scan_like_pdf(tmp_path / "scan.pdf", "Scanned by Office Copier 2026-10-01")
    result = process_pdf(tmp_path / "scan.pdf", TEMPLATES.templates, REGISTRY)
    assert any(n.startswith("page 1 looks scanned") for n in result["notes"])
    assert result["unrecognized_pages"][0]["reason"] == "scanned page: OCR needed"


def test_fingerprint_vocabulary_has_no_pattern_syntax():
    """D7: group names such as (?P<name>...) ended up in the layout fingerprint."""
    template = next(t for t in TEMPLATES.templates if t["template_name"] == "flipkart_seller_tax_invoice.yml")
    words = layout.label_vocabulary(template)
    assert not {"name", "price", "code", "category"} & words and "invoice" in words


def test_poppler_problems_on_a_readable_file_are_reported():
    pdf = need("VIL_bill_7573001478_2026-08-22_statement.pdf")
    notes = process_pdf(pdf, TEMPLATES.templates, REGISTRY)["notes"]
    assert any(n.startswith("pdftotext reported") for n in notes)


# ------------------------------------------------- document types, routing, index, overlap gate

from invoice_engine import overlap  # noqa: E402
from invoice_engine.catalog import Catalog  # noqa: E402
from invoice_engine.document import process_group as _process_group  # noqa: E402
from invoice_engine.export import write_excel  # noqa: E402
from invoice_engine.segment import matching_templates  # noqa: E402

BANK_TYPE = """name: bank_statement
description: Bank account statements
signals: [statement of account, opening balance, closing balance]
min_signals: 2
summary_columns: [issuer, invoice_number, date, amount_opening_balance, amount_closing_balance]
"""
BANK_FIELDS = """document:
  amount_opening_balance: {type: money, role: other}
  amount_credits:         {type: money, role: other}
  amount_debits:          {type: money, role: other}
  amount_closing_balance: {type: money, role: other}
"""
BANK_CHECKS = """document:
  - name: opening + credits - debits = closing
    rule: "amount_opening_balance + amount_credits - amount_debits == amount_closing_balance"
"""
BANK_TEMPLATE = r"""
issuer: Example Bank
keywords: ['EXAMPLE BANK LTD', 'Statement of Account']
required_fields: [issuer, invoice_number, date, amount_closing_balance]
fields:
  document_type: {parser: static, value: bank_statement}
  invoice_number: 'Account No:\s*(\d+)'
  date: 'Statement Date:\s*(\d{2}/\d{2}/\d{4})'
  amount_opening_balance: 'Opening Balance\s+([\d,]+\.\d{2})'
  amount_credits: 'Total Credits\s+([\d,]+\.\d{2})'
  amount_debits: 'Total Debits\s+([\d,]+\.\d{2})'
  amount_closing_balance: 'Closing Balance\s+([\d,]+\.\d{2})'
options: {currency: INR, date_formats: ['%d/%m/%Y']}
"""
BANK_TEXT = """EXAMPLE BANK LTD
Statement of Account
Account No: 001234567
Statement Date: 30/09/2026
Opening Balance      10,000.00
Total Credits         5,250.00
Total Debits          3,100.50
Closing Balance      12,149.50
"""


def _types_with_bank(tmp_path: Path) -> Path:
    """A copy of the real document types plus a bank_statement type: the new type is only files."""
    root = tmp_path / "project"
    shutil.copytree(ROOT / "document_types", root / "document_types")
    shutil.copytree(ROOT / "config", root / "config")
    bank = root / "document_types" / "bank_statement"
    (bank / "templates").mkdir(parents=True)
    (bank / "type.yml").write_text(BANK_TYPE, encoding="utf-8")
    (bank / "fields.yml").write_text(BANK_FIELDS, encoding="utf-8")
    (bank / "checks.yml").write_text(BANK_CHECKS, encoding="utf-8")
    (bank / "templates" / "example_bank_statement.yml").write_text(BANK_TEMPLATE, encoding="utf-8")
    return root


def _run_catalog(text: str, catalog) -> dict:
    groups, _ = split_into_documents([text], catalog)
    assert groups and groups[0]["family"], "no template matched"
    return _process_group(groups[0], ROOT / "test.pdf", catalog)


def test_a_new_document_type_is_only_a_folder(tmp_path: Path):
    catalog = load_catalog(_types_with_bank(tmp_path))
    assert catalog.rejected == {}
    doc = _run_catalog(BANK_TEXT, catalog)
    assert (doc["type"], doc["status"], doc["fields"]["amount_closing_balance"]) == ("bank_statement", "OK", 12149.5)
    wrong = _run_catalog(BANK_TEXT.replace("12,149.50", "12,194.50"), catalog)
    assert wrong["status"] == "FAIL"                      # the type's own check catches a misread


def test_a_field_may_be_defined_once(tmp_path: Path):
    root = _types_with_bank(tmp_path)
    (root / "document_types" / "bank_statement" / "fields.yml").write_text(
        BANK_FIELDS + "  amount: {type: money, role: total}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="already defined"):
        load_catalog(root)


def test_routing_keeps_document_types_apart(tmp_path: Path):
    catalog = load_catalog(_types_with_bank(tmp_path))
    assert [dt.name for dt in catalog.route(BANK_TEXT)] == ["bank_statement"]
    assert {dt.name for dt in catalog.route("nothing that names a document type")} == {dt.name for dt in catalog.types}


def test_keyword_index_finds_the_same_templates_as_a_full_scan():
    for pdf in sorted(ROOT.glob("*.pdf")):
        for text in page_texts(pdf):
            allowed = {dt.name for dt in CATALOG.route(text)}
            full = [t for t in matching_templates(text, CATALOG.templates) if CATALOG.type_of(t).name in allowed]
            assert CATALOG.candidates(text) == full, pdf.name


def test_template_names_are_unique_across_types(tmp_path: Path):
    root = _types_with_bank(tmp_path)
    shared_only = ("issuer: Dup\nkeywords: ['DUP LTD']\nfields:\n  invoice_number: 'No:\\s*(\\S+)'\n"
                   "  date: 'Date:\\s*(\\d{4}-\\d{2}-\\d{2})'\n  amount: 'Total\\s+([\\d.]+)'\n"
                   "options: {currency: INR, date_formats: ['%Y-%m-%d']}\n")
    for type_name in ("bank_statement", "invoice"):
        (root / "document_types" / type_name / "templates" / "dup.yml").write_text(shared_only, encoding="utf-8")
    rejected = load_catalog(root).rejected
    assert list(rejected) == ["invoice/dup.yml"] and "already used" in rejected["invoice/dup.yml"][0]


def test_output_carries_document_type_and_field_pages():
    result = process_pdf(need("VIL_bill_7573001478_2026-08-22_statement.pdf"), CATALOG)
    assert result["format_version"] == 2
    doc = result["documents"][0]
    assert doc["type"] == "invoice"
    assert doc["field_pages"]["invoice_number"] == 1 and doc["field_pages"]["amount_sgst"] == 3


def test_excel_has_one_summary_sheet_per_document_type(tmp_path: Path):
    from openpyxl import load_workbook
    results = [process_pdf(need(name), CATALOG) for name in ("oyo.pdf", "ACK742817410300726.pdf")]
    write_excel(results, tmp_path / "documents.xlsx", CATALOG)
    sheets = load_workbook(tmp_path / "documents.xlsx").sheetnames
    assert sheets[:2] == ["Invoice", "Tax Document"] and "Line Items" in sheets


def test_overlap_gate_finds_clash_overlap_and_missing_sample(tmp_path: Path):
    meesho, oyo = need("ms_invoice_326255343871054720_1.pdf"), need("oyo.pdf")
    derive(tmp_path, "meesho_tax_invoice.yml", "abc_marketplace_invoice.yml", issuer="ABC Marketplace")
    (tmp_path / "com.oyo.invoice.yml").write_text(
        "issuer: OYO\nkeywords: [OYO, Oravel, Stays]\n"
        "fields:\n"
        "  amount: 'Grand Total\\s+Rs (\\d+)'\n"
        "  date: 'Date:\\s(\\d{1,2}/\\d{1,2}/\\d{4})'\n"
        "  invoice_number: '([A-Z0-9]+)\\s+Cash at Hotel'\n"
        "options: {currency: INR, date_formats: ['%d/%m/%Y']}\n", encoding="utf-8")
    catalog = Catalog.single(pool_with(tmp_path), REGISTRY)
    report = overlap.check_samples([meesho, oyo], catalog)
    found = {(f.level, f.where) for f in report.findings}
    assert (overlap.CLASH, f"{meesho.name} pages [1]") in found
    assert (overlap.OVERLAP, f"{oyo.name} pages [1]") in found
    assert (overlap.KEYWORD_SUBSET, "keywords") in found and report.failed()
    assert report.read_by["anthropic_invoice.yml"] == 0 and report.failed("anthropic_invoice.yml")


def test_overlap_gate_is_clean_on_the_real_templates():
    samples = sorted(ROOT.glob("*.pdf"))
    if not samples:
        pytest.skip("no sample PDFs in the project folder")
    report = overlap.check_samples(samples, CATALOG)
    assert [f for f in report.findings] == [], [(f.level, f.where, f.message) for f in report.findings]
    # every template whose sample PDF is in the folder must read it (PDFs not copied here are skipped)
    expected = [json.loads(f.read_text(encoding="utf-8")) for f in EXPECTED]
    covered = {d["template"] for e in expected if (ROOT / e["source_file"]).exists() for d in e["documents"]}
    assert all(report.read_by.get(t, 0) > 0 for t in covered), [t for t in covered if not report.read_by.get(t)]


# --- optional_fields ----------------------------------------------------------------------------

def test_optional_field_missing_does_not_make_check(tmp_path: Path):
    """A field listed in optional_fields may be absent without a warning; other absent fields
    still warn."""
    extra = "  customer_number: 'Customer Ref:\\s*(\\S+)'\n  amount_cess:"
    plain = CESS_TEMPLATE.replace("  amount_cess:", extra, 1)
    write_template(tmp_path, "acme.yml", plain)
    doc = run_text(CESS_TEXT, load_templates(tmp_path, REGISTRY).templates)
    assert any("found nothing: customer_number" in i["message"] for i in doc["issues"])

    optional = plain.replace("required_fields:", "optional_fields: [customer_number]\nrequired_fields:", 1)
    write_template(tmp_path, "acme.yml", optional)
    doc = run_text(CESS_TEXT, load_templates(tmp_path, REGISTRY).templates)
    assert not any("customer_number" in i["message"] for i in doc["issues"])
    assert doc["status"] == "OK"


def test_optional_fields_lint():
    from invoice_engine.lint import lint_template
    import yaml as _yaml
    spec = _yaml.safe_load(CESS_TEMPLATE)
    spec["template_name"] = "acme.yml"
    spec["optional_fields"] = ["no_such_field", "invoice_number"]
    errors, _ = lint_template(spec, REGISTRY)
    assert any("'no_such_field' is not one of" in e for e in errors)
    assert any("'invoice_number' is also in required_fields" in e for e in errors)


# --- layout fingerprints built by default ---------------------------------------------------------

def test_build_layouts_learns_only_from_documents_the_template_won():
    from invoice_engine.fingerprints import build_layouts
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    results = {r.template: r for r in build_layouts(CATALOG, [pdf], write=False)}
    assert results["zomato_restaurant_invoice.yml"].samples == ["Invoice_8626225731.pdf pages 1 (OK)"]
    assert results["zomato_platform_fee_invoice.yml"].samples == ["Invoice_8626225731.pdf pages 2 (OK)"]
    assert results["zomato_restaurant_invoice.yml"].labels > 0
    assert not results["shopify_bill.yml"].samples          # not a sample of other layouts
    assert not any(r.written for r in results.values())   # write=False changes no file


def test_every_template_has_a_layout_block():
    missing = [t["template_name"] for t in CATALOG.templates if not t.get("layout")]
    assert not missing, ("templates without a layout block; run "
                         "python process_invoice.py --fingerprint <sample folder>: " + ", ".join(missing))


# --- a number the document does not print stays empty -----------------------------------------

def test_empty_typed_number_is_none_not_zero(tmp_path: Path):
    """invoice2data turns an empty int/float into 0; the engine keeps it empty (compat.py)."""
    from invoice2data.extract.invoice_template import InvoiceTemplate
    template = load_templates(INVOICE_TEMPLATES, REGISTRY).templates[0]
    assert isinstance(template, InvoiceTemplate)
    assert template.coerce_type("", "int") is None
    assert template.coerce_type("  ", "float") is None
    assert template.coerce_type("12", "int") == 12


# --- standard form of page text and the stress test ---------------------------------------------

def test_canonical_text_removes_print_differences():
    from invoice_engine.segment import canonical
    a = "Invoice No.: A1\nTotal:  5.00\n\nThanks\n"
    b = "      Invoice No. : A1   \n      Total:  5.00\n\n\n\n      Thanks\n"
    assert canonical(a) == canonical(b) == a
    # padding that lines up a column of colons moves after the colon; columns stay in place
    assert canonical("Terms      : Due      Place: X") == "Terms:       Due      Place: X"
    assert canonical("Time 11 : 04 and https://x") == "Time 11 : 04 and https://x"   # not a label


def test_stress_test_reports_a_fragile_pattern(tmp_path: Path):
    from invoice_engine.stress import compare, REWRITES
    change = REWRITES["short decimals"][1]
    original = {"status": "OK", "fields": {"amount": 212.1, "note": "Total 5.00"}, "lines": []}
    same = {"status": "OK", "fields": {"amount": 212.1, "note": "Total 5"}, "lines": []}
    wrong = {"status": "OK", "fields": {"amount": 21.0, "note": "Total 5"}, "lines": []}
    assert compare(original, same, change) == ([], [])                 # rewritten text counts as same
    assert compare(original, wrong, change) == ([], ["amount"])
    assert compare(original, {"status": "OK", "fields": {}, "lines": []}, change)[0] == ["amount", "note"]


def test_templates_survive_print_variations_on_their_samples():
    """No rewrite may change a value silently (WRONG); the zomato templates must read every
    rewrite exactly as the original."""
    from invoice_engine.stress import stress
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    reports = stress(CATALOG, [pdf])
    assert not [o for r in reports for o in r.outcomes if o.result == "WRONG"]
    for name in ("zomato_restaurant_invoice.yml",):
        report = next(r for r in reports if r.template == name)
        assert report.counts()["same"] == len(report.outcomes)


# --- automatic template checks in the pipeline ---------------------------------------------------

def _pipeline(tmp_path, catalog=None):
    from invoice_engine.pipeline import Pipeline
    from invoice_engine.samples import SampleStore
    return Pipeline(catalog or CATALOG, SampleStore(tmp_path / "samples"), write_layouts=False)


def test_pipeline_admits_on_first_document_and_stores_it(tmp_path: Path):
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    pipe = _pipeline(tmp_path)
    result = pipe.process(pdf)
    assert [d["status"] for d in result["documents"]] == ["OK", "OK"]
    assert pipe.state["zomato_restaurant_invoice.yml"] == "admitted"
    assert len(pipe.store.data["samples"]) == 2
    # nothing changed: a second pipeline does no checks and stores nothing new
    again = _pipeline(tmp_path)
    again.process(pdf)
    assert not again.events and len(again.store.data["samples"]) == 2


def test_changed_template_with_different_values_is_held(tmp_path: Path):
    """A template edit that keeps the sums right but reads another value is not trusted."""
    import shutil
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    _pipeline(tmp_path).process(pdf)                        # stores the verified values
    types = tmp_path / "document_types"
    shutil.copytree(ROOT / "document_types", types)
    shutil.copytree(ROOT / "config", tmp_path / "config")
    target = types / "invoice" / "templates" / "zomato_restaurant_invoice.yml"
    text = target.read_text(encoding="utf-8")
    old = "  seller_name: 'Legal Entity Name\\s*:\\s*(.+?)\\s*\\n'"
    assert old in text
    target.write_text(text.replace(old, "  seller_name: 'Restaurant Name\\s*:\\s*(.+?)\\s*\\n'"), encoding="utf-8")
    pipe = _pipeline(tmp_path, load_catalog(types, tmp_path / "config"))
    pipe.prepare()
    assert pipe.state["zomato_restaurant_invoice.yml"] == "held"
    assert any("seller_name" in p for e in pipe.events for p in e.problems)
    assert "zomato_restaurant_invoice.yml" not in {t["template_name"] for t in pipe.catalog.templates}


def test_new_template_that_reads_another_layout_is_held_before_use(tmp_path: Path):
    import shutil
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    _pipeline(tmp_path).process(pdf)
    types = tmp_path / "document_types"
    shutil.copytree(ROOT / "document_types", types)
    shutil.copytree(ROOT / "config", tmp_path / "config")
    source = (types / "invoice" / "templates" / "zomato_platform_fee_invoice.yml").read_text(encoding="utf-8")
    copy = source.split("\nlayout:")[0].replace("issuer: Zomato (Eternal Limited)", "issuer: Copy Co")
    (types / "invoice" / "templates" / "zz_copy.yml").write_text(copy, encoding="utf-8")
    pipe = _pipeline(tmp_path, load_catalog(types, tmp_path / "config"))
    result = pipe.process(pdf)
    assert pipe.state["zz_copy.yml"] == "held"
    assert [d["status"] for d in result["documents"]] == ["OK", "OK"]   # the real template still reads it


# --- AI template author (with a stand-in for the AI service) -------------------------------------

class FakeAI:
    """Answers the classify and write prompts from prepared lists, in order."""

    def __init__(self, classify: dict | None, drafts: list[str]):
        self.classify, self.drafts, self.prompts = classify, list(drafts), []

    def complete(self, system: str, prompt: str) -> str:
        if prompt.startswith("Read this document and copy the value"):     # second opinion: no values
            self.opinions = getattr(self, "opinions", 0) + 1
            return "```json\n{}\n```"
        self.prompts.append(prompt)
        if prompt.startswith("Classify"):
            return "```json\n" + json.dumps(self.classify) + "\n```"
        return "```yaml\n" + (self.drafts.pop(0) if self.drafts else "not: [valid") + "\n```"


def _project(tmp_path: Path, drop: list[str] = (), drop_types: list[str] = ()) -> tuple[Path, Path]:
    """A copy of document_types and config without some templates / types."""
    types = tmp_path / "document_types"
    shutil.copytree(ROOT / "document_types", types)
    shutil.copytree(ROOT / "config", tmp_path / "config")
    for name in drop:
        for path in types.glob(f"*/templates/{name}"):
            path.unlink()
    for name in drop_types:
        shutil.rmtree(types / name)
    return types, tmp_path / "config"


@pytest.fixture(autouse=True)
def _hand_written_drafts(request, monkeypatch):
    """The fake AI of the test_ai_* tests answers with templates written by hand before the GST
    coverage rule (most read IGST only); those tests check other parts of the author."""
    if request.node.name.startswith("test_ai_"):
        monkeypatch.setattr("invoice_engine.ai_author._gst_coverage", lambda data: [])


def _ai_pipeline(tmp_path: Path, types: Path, config: Path, fake: FakeAI):
    from invoice_engine.ai_author import TemplateAuthor
    from invoice_engine.ai_provider import AISettings
    from invoice_engine.pipeline import Pipeline
    from invoice_engine.samples import SampleStore
    store = SampleStore(tmp_path / "samples")
    author = TemplateAuthor(fake, AISettings(enabled=True, model="fake", max_attempts=3), types, config, store)
    return Pipeline(load_catalog(types, config), store, write_layouts=False, author=author,
                    reload=lambda: load_catalog(types, config))


def _body(name: str) -> str:
    path = next((ROOT / "document_types").glob(f"*/templates/{name}"))
    return path.read_text(encoding="utf-8").split("\nlayout:")[0]


def test_ai_writes_a_template_after_a_correction_round(tmp_path: Path):
    pdf = need("TI26-2701230.pdf")
    types, config = _project(tmp_path, drop=["mtalkz_tax_invoice.yml"])
    good = _body("mtalkz_tax_invoice.yml")
    wrong = good.replace("Sub Total\\s+([\\d,]+(?:\\.\\d{1,2})?)", "IGST\\d*\\s*\\(([\\d.]+)%\\)")   # untaxed = the rate
    assert wrong != good
    fake = FakeAI({"document_type": "invoice", "is_new_type": False, "issuer": "Mtalkz",
                   "variant_of": None, "template_name": "mtalkz_invoice"}, [wrong, good])
    pipe = _ai_pipeline(tmp_path, types, config, fake)
    result = pipe.process(pdf)
    assert [d["status"] for d in result["documents"]] == ["OK"]
    assert result["documents"][0]["template"] == "mtalkz_invoice.yml"
    ai = result["ai"][0]
    assert (ai["outcome"], ai["attempts"]) == ("admitted", 2)
    assert "rejected" in fake.prompts[-1]                        # the second draft saw the problems
    assert (types / "invoice" / "templates" / "mtalkz_invoice.yml").is_file()
    # the same document is not sent to the AI again
    again = _ai_pipeline(tmp_path, types, config, FakeAI(None, []))
    assert again.process(pdf)["documents"][0]["status"] == "OK"


def test_ai_creates_a_new_document_type(tmp_path: Path):
    pdf = need("ACK735221680300726.pdf")
    types, config = _project(tmp_path, drop_types=["tax_document"])
    fields = yaml.safe_load((ROOT / "document_types" / "tax_document" / "fields.yml").read_text(encoding="utf-8"))
    draft = yaml.safe_load(_body("income_tax_itr_acknowledgement.yml"))
    draft["new_fields"] = fields["document"]
    fake = FakeAI({"document_type": "tax_document", "is_new_type": True, "issuer": "Income Tax Department",
                   "variant_of": None, "template_name": "itr_acknowledgement",
                   "new_type": {"name": "tax_document", "description": "Income tax filings",
                                "signals": ["income tax return", "assessment year", "acknowledgement number"],
                                "min_signals": 2, "summary_columns": ["issuer", "invoice_number", "date"]}},
                  [yaml.safe_dump(draft, sort_keys=False)])
    pipe = _ai_pipeline(tmp_path, types, config, fake)
    result = pipe.process(pdf)
    assert result["ai"][0]["outcome"] == "admitted" and result["ai"][0]["new_type"]
    assert (types / "tax_document" / "type.yml").is_file()
    assert "taxpayer_pan" in (types / "tax_document" / "fields.yml").read_text(encoding="utf-8")
    assert [(d["type"], d["status"]) for d in result["documents"]] == [("tax_document", "OK")]


def test_ai_failure_leaves_nothing_behind(tmp_path: Path):
    pdf = need("ACK735221680300726.pdf")
    types, config = _project(tmp_path, drop_types=["tax_document"])
    before = sorted(str(p.relative_to(types)) for p in types.rglob("*"))
    fake = FakeAI({"document_type": "tax_document", "is_new_type": True, "issuer": "Income Tax Department",
                   "template_name": "itr", "new_type": {"name": "tax_document", "description": "x",
                                                         "signals": ["income tax return"]}},
                  ["issuer: x\nkeywords: ['NOT ON THE PAGE']\nfields: {}"] * 3)
    pipe = _ai_pipeline(tmp_path, types, config, fake)
    result = pipe.process(pdf)
    assert result["ai"][0]["outcome"] == "failed" and result["ai"][0]["attempts"] == 3
    assert sorted(str(p.relative_to(types)) for p in types.rglob("*")) == before   # type folder removed
    assert (tmp_path / "samples" / "ai_drafts" / "itr.yml").is_file()               # last draft kept
    assert result["unrecognized_pages"]


def test_ai_writes_a_new_version_when_a_template_fails_on_a_variant(tmp_path: Path):
    first, variant = need("Invoice_8626225731.pdf"), need("Invoice_8261790801.pdf")
    types, config = _project(tmp_path)
    path = types / "invoice" / "templates" / "zomato_platform_fee_invoice.yml"
    current = _body("zomato_platform_fee_invoice.yml")
    broken = current.replace("(?:\\.\\d{1,3})?", "\\.\\d{2}")                 # cannot read 17.582
    assert broken != current
    path.write_text(broken, encoding="utf-8")
    fake = FakeAI(None, [current])
    pipe = _ai_pipeline(tmp_path, types, config, fake)
    assert [d["status"] for d in pipe.process(first)["documents"]] == ["OK", "OK"]     # stored as samples
    result = pipe.process(variant)
    ai = result["ai"][0]
    assert (ai["action"], ai["outcome"], ai["template"]) == ("new version", "admitted", "zomato_platform_fee_invoice_v2.yml")
    assert [d["status"] for d in result["documents"]] == ["OK", "OK"]
    assert result["documents"][1]["template"] == "zomato_platform_fee_invoice_v2.yml"
    assert "THE CURRENT TEMPLATE FAILS" in fake.prompts[0]


def test_ai_draft_in_a_wrong_shape_is_rejected_not_a_crash(tmp_path: Path):
    """A model wrote `fields:` as a list; that stopped the whole file (AttributeError)."""
    pdf = need("TI26-2701230.pdf")
    types, config = _project(tmp_path, drop=["mtalkz_tax_invoice.yml"])
    bad = ("issuer: x\nkeywords: ['Mtalkz Mobility Services Private Limited']\n"
           "fields:\n  - invoice_number: 'Invoice No\\.\\s*:\\s*(\\S+)'\n")
    fake = FakeAI({"document_type": "invoice", "issuer": "Mtalkz", "template_name": "mtalkz"},
                  [bad, "options: [a, b]\nkeywords: x", _body("mtalkz_tax_invoice.yml")])
    result = _ai_pipeline(tmp_path, types, config, fake).process(pdf)
    assert result["ai"][0]["outcome"] == "admitted" and result["ai"][0]["attempts"] == 3
    assert "fields must be a mapping" in fake.prompts[2]           # the model was told what was wrong


def test_one_unreadable_template_file_does_not_stop_the_others(tmp_path: Path):
    types, config = _project(tmp_path)
    (types / "invoice" / "templates" / "zz_bad.yml").write_text(
        "issuer: x\nkeywords: ['A']\noptions: [currency, INR]\nfields: {}\n", encoding="utf-8")
    catalog = load_catalog(types, config)
    assert len(catalog.templates) == len(CATALOG.templates)
    assert "invoice/zz_bad.yml" in catalog.rejected


def test_ai_keywords_are_fitted_to_the_page_text():
    """Models write column text with single spaces; such keywords are fitted to the page."""
    from invoice_engine.ai_author import _fit_keywords
    text = "Booked From                                  To\nTotal Fare (all inclusive)   10,017.20\n"
    data = {"keywords": ["Booked From To", "Total Fare (all inclusive)", "Not On Page"],
            "exclude_keywords": ["Booked From"]}
    problems, notes = _fit_keywords(data, text)
    assert data["keywords"] == ["Booked From", "Total Fare (all inclusive)"]   # one column only
    assert data["exclude_keywords"] == []                                   # it was on this page
    assert problems == ["keyword 'Not On Page' is not in the document text (not even with other spacing): "
                        "choose a text that is printed on the page, on one line"]


# --- AI provider: the ways a model can answer ---------------------------------------------------

def _serve(answer: dict):
    """A local server with the OpenAI chat API that gives one fixed answer."""
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_POST(self):
            self.rfile.read(int(self.headers["Content-Length"]))
            data = json.dumps(answer).encode()
            self.send_response(200)
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

    server = HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


@pytest.mark.parametrize("message,finish,expected", [
    ({"content": "ok"}, "stop", "ok"),
    ({"content": "<think>let me see {not json}</think>\n{\"a\": 1}"}, "stop", '{"a": 1}'),
    ({"content": None, "reasoning_content": "{\"a\": 1}"}, "stop", '{"a": 1}'),
    ({"content": ""}, "length", "cut off at ai.max_tokens"),
    ({"content": ""}, "stop", "empty answer"),
])
def test_ai_provider_reads_every_answer_shape(message, finish, expected, monkeypatch):
    from invoice_engine.ai_provider import AIError, AISettings, OpenAICompatibleProvider
    server = _serve({"choices": [{"message": message, "finish_reason": finish}]})
    try:
        monkeypatch.setenv("TEST_AI_KEY", "x")
        provider = OpenAICompatibleProvider(AISettings(base_url=f"http://127.0.0.1:{server.server_port}/v1",
                                                       model="m", api_key_env="TEST_AI_KEY"))
        try:
            assert provider.complete("s", "p") == expected
        except AIError as exc:
            assert expected in str(exc)
    finally:
        server.shutdown()


def test_ai_classification_falls_back_to_routing(tmp_path: Path):
    """Two unreadable classification answers: the signal words pick the type, drafting goes on."""
    pdf = need("TI26-2701230.pdf")
    types, config = _project(tmp_path, drop=["mtalkz_tax_invoice.yml"])

    class Prose(FakeAI):
        def complete(self, system, prompt):
            self.prompts.append(prompt)
            if prompt.startswith("Classify"):
                return "This looks like a tax invoice from a messaging company."
            return "```yaml\n" + _body("mtalkz_tax_invoice.yml") + "\n```"

    fake = Prose(None, [])
    result = _ai_pipeline(tmp_path, types, config, fake).process(pdf)
    assert result["ai"][0]["outcome"] == "admitted"
    assert sum(p.startswith("Classify") for p in fake.prompts) == 2
    assert "could not be read" in fake.prompts[1]


def test_ai_sum_error_lists_the_unread_rows_and_sum_checks_stay_on(tmp_path: Path):
    pdf = need("TI26-2701230.pdf")
    types, config = _project(tmp_path, drop=["mtalkz_tax_invoice.yml"])
    good = _body("mtalkz_tax_invoice.yml")
    # rows with a rate below 0.1 are not read by this draft: the DLT row (0.025) is missed
    weak = good.replace("(?P<price_unit>[\\d,]*\\.\\d+)", "(?P<price_unit>[\\d,]*\\.\\d\\d?(?!\\d))")
    sneaky = weak.replace("checks:\n", "checks:\n  disable: [sum of line price_subtotal vs amount_untaxed]\n")
    assert weak != good and sneaky != weak
    fake = FakeAI({"document_type": "invoice", "issuer": "Mtalkz", "template_name": "mtalkz"}, [weak, sneaky, good])
    result = _ai_pipeline(tmp_path, types, config, fake).process(pdf)
    assert result["ai"][0]["outcome"] == "admitted" and result["ai"][0]["attempts"] == 3
    assert "DLT Scrubbing Charges" in fake.prompts[2]          # the missed row, shown to the model
    assert "cannot be disabled" in fake.prompts[3]


def test_ai_mechanical_slips_are_repaired_and_broken_optional_fields_dropped(tmp_path: Path):
    """A missing (?m) is added; an optional field with a broken pattern is removed instead of
    rejecting a template whose required fields and arithmetic are right."""
    pdf = need("TI26-2701230.pdf")
    types, config = _project(tmp_path, drop=["mtalkz_tax_invoice.yml"])
    draft = _body("mtalkz_tax_invoice.yml")
    draft = draft.replace("payment_terms: '(?m)^Terms", "payment_terms: '^Terms")           # no (?m)
    draft = draft.replace("  seller_email: '\\d{10}:\\s*(\\S+@\\S+)'",
                          "  seller_email: '(\\d{10}):\\s*(\\S+@\\S+)'")                     # 2 groups
    assert "payment_terms: '^Terms" in draft and "(\\d{10}):" in draft
    fake = FakeAI({"document_type": "invoice", "issuer": "Mtalkz", "template_name": "mtalkz"}, [draft])
    result = _ai_pipeline(tmp_path, types, config, fake).process(pdf)
    assert result["ai"][0]["outcome"] == "admitted" and result["ai"][0]["attempts"] == 1
    written = (types / "invoice" / "templates" / "mtalkz.yml").read_text(encoding="utf-8")
    assert "(?m)^Terms" in written and "seller_email" not in written
    assert result["documents"][0]["fields"]["payment_terms"] == "Due on Receipt"


def test_ai_correction_rounds_include_the_document_text(tmp_path: Path):
    """The model was blind after the first draft ("I don't have the document text")."""
    pdf = need("TI26-2701230.pdf")
    types, config = _project(tmp_path, drop=["mtalkz_tax_invoice.yml"])
    good = _body("mtalkz_tax_invoice.yml")
    fake = FakeAI({"document_type": "invoice", "issuer": "Mtalkz", "template_name": "mtalkz"},
                  ["issuer: x\nkeywords: ['Mtalkz Mobility Services Private Limited']\nfields: {}\n", good])
    _ai_pipeline(tmp_path, types, config, fake).process(pdf)
    assert "DOCUMENT TEXT" in fake.prompts[2] and "OTP Sms" in fake.prompts[2]


def test_total_rounding_compares_the_rounded_sum_exactly(tmp_path: Path):
    """IRCTC rounds the total up to whole rupees: 518.15 -> 519, 1378.60 -> 1379."""
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    from invoice_engine.validate import check_document
    base = {"issuer": "x", "document_type": "t", "invoice_number": "1", "date": "2026-01-01", "seller_name": "x"}
    template["fields"] = {k: "x" for k in list(base) + ["amount"]}

    def sum_error(amount, rows, **rounding):
        t = dict(template, **rounding)
        issues, _ = check_document(dict(base, amount=amount), [{"name": str(i), "price_total": v} for i, v in enumerate(rows)],
                                   [], t, REGISTRY)
        return any("sum of line price_total vs amount" in m for _, m in issues)

    assert sum_error(519.0, [500.0, 18.15])                                              # no rounding declared
    assert not sum_error(519.0, [500.0, 18.15], total_rounding=1, total_rounding_mode="up")
    assert sum_error(519.0, [500.0, 18.15], total_rounding=1, total_rounding_mode="nearest")   # 518.15 -> 518
    assert sum_error(520.0, [500.0, 18.15], total_rounding=1, total_rounding_mode="up")        # off by more
    assert not sum_error(1379.0, [1340.0, 38.6], total_rounding=1)                             # nearest
    assert sum_error(518.5, [500.0, 18.15], total_rounding=1, total_rounding_mode="up")        # not a whole total


def test_a_document_number_without_digits_is_flagged():
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    from invoice_engine.validate import check_document
    template["fields"] = {k: "x" for k in ("issuer", "document_type", "invoice_number", "date", "amount")}
    fields = {"issuer": "x", "document_type": "t", "invoice_number": "Train", "date": "2026-01-01", "amount": 5.0}
    issues, _ = check_document(fields, [], [], template, REGISTRY)
    assert any("has no digit" in m for _, m in issues)
    issues, _ = check_document(dict(fields, invoice_number="100006522561796"), [], [], template, REGISTRY)
    assert not any("has no digit" in m for _, m in issues)


def test_ai_is_told_about_pages_left_out_of_the_document():
    from invoice_engine.ai_author import _leftover_pages
    text = ("=== page 1 ===\nElectronic Reservation Slip\nTotal Fare:  Rs. 1379/-\n"
            "=== page 2 ===\nSupplier Information:\nSAC Code:      GSTIN: N.A\nRecipient Information:\n")
    problem = _leftover_pages(text, [1, 2], {"pages": [1]})
    assert problem and "continuation_keywords" in problem[0] and "'Supplier Information:'" in problem[0]
    assert _leftover_pages(text, [1, 2], {"pages": [1, 2]}) == []


def test_joined_table_rows_get_a_clear_error():
    from invoice_engine.compat import JoinedRowsError, install
    from invoice2data.extract.invoice_template import InvoiceTemplate
    install()
    template = InvoiceTemplate({"issuer": "x", "keywords": ["x"], "fields": {}, "template_name": "t.yml",
                                "options": {"decimal_separator": "."}})
    with pytest.raises(JoinedRowsError, match="several rows"):
        template.coerce_type("1340\n17.7\n20", "float")
    assert template.coerce_type("17.7", "float") == 17.7


def test_ai_draft_with_two_amounts_from_one_pattern_is_rejected():
    from invoice_engine.ai_author import _duplicate_fields
    data = {"fields": {"amount": "Total Fare:\\s+(\\d+)", "amount_untaxed": "Total Fare:\\s+(\\d+)",
                       "invoice_number": "ID:\\s*(\\d+)"}}
    assert _duplicate_fields(data) and "amount_untaxed" in _duplicate_fields(data)[0]
    assert _duplicate_fields({"fields": {"amount": "A(\\d+)", "amount_tax": "B(\\d+)"}}) == []


def test_a_cut_off_answer_is_not_taken_from_the_thinking():
    from invoice_engine.ai_provider import AIError, AISettings, OpenAICompatibleProvider
    provider = OpenAICompatibleProvider(AISettings(model="m"))
    provider._request = lambda path, body=None: {"choices": [{"finish_reason": "length", "message": {
        "content": "", "reasoning_content": "The user wants me to fix ... to"}}]}
    with pytest.raises(AIError, match="still thinking"):
        provider.complete("s", "p")


def test_total_rounding_written_under_options_still_counts():
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    from invoice_engine.validate import check_document
    base = {"issuer": "x", "document_type": "t", "invoice_number": "1", "date": "2026-01-01", "seller_name": "x"}
    template["fields"] = {k: "x" for k in list(base) + ["amount"]}
    template["options"] = dict(template.get("options") or {}, total_rounding=1, total_rounding_mode="up")
    issues, _ = check_document(dict(base, amount=519.0), [{"name": "a", "price_total": 500.0},
                                                          {"name": "b", "price_total": 18.15}], [], template, REGISTRY)
    assert not any("sum of line" in m for _, m in issues)


def test_ai_draft_rounding_is_moved_out_of_options_and_empty_fields_are_reported():
    from invoice_engine.ai_author import _empty_fields, _repair_mechanical
    data = {"options": {"currency": "INR", "total_rounding": 1, "total_rounding_mode": "up"},
            "fields": {"amount": "Total:\\s+(\\d+)", "travel_class": "Train No/Name\\s+(\\S+)",
                       "document_type": {"parser": "static", "value": "invoice"}}}
    notes = _repair_mechanical(data)
    assert data["total_rounding"] == 1 and data["total_rounding_mode"] == "up"
    assert "total_rounding" not in data["options"] and notes
    problem = _empty_fields(data, {"fields": {"amount": 519.0, "document_type": "invoice"}})
    assert problem and "travel_class" in problem[0] and "amount" not in problem[0].split(":")[1]


# --- tables read by their column headings (parser: table) -------------------------------------

def _table(table_text: str, **extra):
    from invoice2data.extract.invoice_template import InvoiceTemplate
    from invoice_engine import table
    settings = {"start": r"(?m)^Product\s+Title", "end": r"(?m)^\s+Total\s+\d",
                "columns": {"category": "^Product", "name": "Title", "qty": "Qty", "gross_amount": "Gross",
                            "discount": "Discount", "price_subtotal": "Taxable", "igst_amount": "IGST",
                            "cgst_amount": "CGST", "sgst_amount": "SGST", "cess_amount": "CESS",
                            "price_total": "Total"},
                "types": {k: "float" for k in ("gross_amount", "discount", "price_subtotal", "igst_amount",
                                               "cgst_amount", "sgst_amount", "cess_amount", "price_total")}
                         | {"qty": "int"}}
    settings.update(extra)
    template = InvoiceTemplate({"issuer": "x", "keywords": ["x"], "fields": {}, "template_name": "t.yml",
                                "options": {"decimal_separator": "."}})
    return table.parse(template, "lines", settings, table_text)


def test_table_reads_one_or_two_tax_columns_with_the_same_settings():
    igst = (
        "Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n"
        "\n"
        "Shoes          RED TAPE Sneakers for       1   785.00    -135.00    619.05   30.95    650.00\n"
        "               Men | Premium\n"
        "               Shipping Charges            1    40.00     -40.00      0.00    0.00      0.00\n"
        "                          Total            2   825.00    -175.00    619.05   30.95    650.00\n")
    rows = _table(igst)
    assert [r["price_total"] for r in rows] == [650.0, 0.0]
    assert rows[0]["igst_amount"] == 30.95 and rows[0]["name"] == "RED TAPE Sneakers for\nMen | Premium"
    assert rows[0]["category"] == "Shoes" and rows[1]["name"] == "Shipping Charges" and "category" not in rows[1]
    cgst_sgst = (
        "Product        Title                     Qty   Gross    Discount   Taxable    CGST   SGST     Total\n"
        "\n"
        "Earbuds        Boult Audio X10             1   899.00     -59.00    711.86   64.07   64.07    840.00\n"
        "                          Total            1   899.00     -59.00    711.86   64.07   64.07    840.00\n")
    row = _table(cgst_sgst)[0]
    assert (row["cgst_amount"], row["sgst_amount"], row["price_total"]) == (64.07, 64.07, 840.0)
    assert "igst_amount" not in row


def test_table_handles_a_missing_column_and_numbers_inside_descriptions():
    text = (
        "Product        Title                            Gross    Taxable    IGST    CESS     Total\n"
        "\n"
        "Snacks         Chips Pack of 2                  100.00     84.75   15.25    0.00    100.00\n"
        "               Pack of 2 x 50 g\n"
        "                          Total                 100.00     84.75   15.25    0.00    100.00\n")
    row = _table(text, row=["price_total"])[0]
    assert "qty" not in row and row["cess_amount"] == 0.0 and row["price_subtotal"] == 84.75
    assert row["name"] == "Chips Pack of 2\nPack of 2 x 50 g"      # the 2 and 50 stay in the description


def test_table_text_before_the_amounts_row_and_a_row_end():
    text = (
        "Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n"
        "\n"
        "Helmet Black   Full face\n"
        "                                           1   404.00     -50.00    300.00   54.00    354.00\n"
        "IMEI: []\n"
        "Shipping\n"
        "                                           1     0.00       0.00      0.00    0.00      0.00\n"
        "                          Total            2   404.00     -50.00    300.00   54.00    354.00\n")
    rows = _table(text, text="after", row_end="^IMEI")
    # text is placed by its own column: "Shipping" is printed under Product
    assert len(rows) == 2 and rows[1]["category"] == "Shipping" and "name" not in rows[1]
    assert rows[0]["category"] == "Helmet Black\nIMEI: []" and rows[0]["name"] == "Full face"
    rows = _table(text, text="before")                       # without row_end, text after a row
    assert rows[1]["category"] == "IMEI: []\nShipping"       # goes to the next row


def test_table_two_values_for_one_column_are_an_error_not_a_choice():
    from invoice_engine.compat import JoinedRowsError
    text = (
        "Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n"
        "\n"
        "Shoes          Sneakers                    1   785.00    -135.00    619.05   30.95    650.00\n"
        "                                                                             31.00\n"
        "                          Total            1   785.00    -135.00    619.05   30.95    650.00\n")
    with pytest.raises(JoinedRowsError):
        _table(text)


def test_table_a_column_the_template_does_not_list_is_an_error_never_another_column():
    from invoice_engine.table import UnknownColumnError
    no_room = ("Product        Title                     Qty   Gross    Discount   Taxable    IGST     TCS     Total\n\n"
               "Shoes          Sneakers                    1   785.00    -135.00    619.05   30.95    3.10    653.10\n"
               "                          Total            1   785.00    -135.00    619.05   30.95    3.10    653.10\n")
    with pytest.raises(UnknownColumnError, match="TCS"):
        _table(no_room)
    # an empty listed column (CESS) next to the unknown one must not take its value
    next_to_empty = ("Product        Title                     Qty   Gross    Discount   Taxable    IGST     CESS    TCS     Total\n\n"
                     "Shoes          Sneakers                    1   785.00    -135.00    619.05   30.95            3.10   653.10\n"
                     "                          Total            1   785.00    -135.00    619.05   30.95            3.10   653.10\n")
    with pytest.raises(UnknownColumnError, match="'TCS'.*cess_amount"):
        _table(next_to_empty)


# --- table: the cases a real document can bring (each was a silent or a needless failure) --------

_HDR = "Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n"
_END = "                          Total            2   825.00    -175.00    619.05   30.95    650.00\n"


def test_table_continues_after_a_page_break_without_taking_the_footer_or_the_repeated_header():
    text = (_HDR + "\n"
            "Shoes          RED TAPE Sneakers           1   785.00    -135.00    619.05   30.95    650.00\n"
            "               for Men\n"
            "\n"
            "  Returns Policy: At Flipkart we try to deliver perfectly each and every time. But in the off-chance\n"
            "                                                                          page 1 of 2\n"
            # page 2: the header again, printed a little further right
            "   Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n"
            "\n"
            "                  Shipping Charges            1    40.00     -40.00      0.00    0.00      0.00\n"
            + _END)
    rows = _table(text, start=r"(?m)^\s*Product\s+Title")
    assert [r["price_total"] for r in rows] == [650.0, 0.0]
    assert rows[0]["name"] == "RED TAPE Sneakers\nfor Men" and rows[1]["name"] == "Shipping Charges"


def test_table_two_columns_with_the_same_heading_are_a_template_error():
    from invoice2data.exceptions import TemplateSyntaxError
    text = (_HDR + "\n"
            "Shoes          Sneakers                    1   785.00    -135.00    619.05   30.95    650.00\n" + _END)
    with pytest.raises(TemplateSyntaxError, match="same heading"):
        _table(text, columns={"name": "Title", "qty": "Qty", "price_subtotal": "Taxable", "price_total": "Ta"})


def test_table_number_formats():
    text = ("Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n\n"
            "Shoes          Sneakers                    1   785.00    (135.00)   619.05   30.95    650.00\n"
            "Socks          Cotton                      1   100.00     10.00-     84.75   15.25     90.00\n"
            "Belt           Leather                     1   $100.00        -      84.75   15.25    100.00\n"
            "Cap            Blue                        2 Nos  50.00       NA     42.37    7.63     50.00\n"
            + _END)
    rows = _table(text)
    assert rows[0]["discount"] == -135.0 and rows[1]["discount"] == -10.0
    assert rows[2]["gross_amount"] == 100.0 and "discount" not in rows[2] or rows[2]["discount"] is None
    assert rows[3]["qty"] == 2 and rows[3]["name"] == "Blue"


def test_table_decimal_comma():
    from invoice2data.extract.invoice_template import InvoiceTemplate
    from invoice_engine import table
    text = ("Product        Title                     Qty   Gross    Discount   Taxable    IGST     Total\n\n"
            "Boek           Roman                       1   1.234,56      0,00   1.020,30  214,26  1.234,56\n"
            "                          Total            1   1.234,56      0,00   1.020,30  214,26  1.234,56\n")
    template = InvoiceTemplate({"issuer": "x", "keywords": ["x"], "fields": {}, "template_name": "t.yml",
                                "options": {"decimal_separator": ","}})
    settings = {"start": r"(?m)^Product", "end": r"(?m)^\s+Total\s+\d",
                "columns": {"category": "^Product", "name": "Title", "qty": "Qty", "gross_amount": "Gross",
                            "discount": "Discount", "price_subtotal": "Taxable", "igst_amount": "IGST",
                            "price_total": "Total"},
                "types": {"qty": "int", "gross_amount": "float", "discount": "float", "price_subtotal": "float",
                          "igst_amount": "float", "price_total": "float"}}
    row = table.parse(template, "lines", settings, text)[0]
    assert row["price_total"] == 1234.56 and row["igst_amount"] == 214.26


def test_table_rate_in_brackets_stays_positive():
    from invoice_engine.table import clean_number
    assert clean_number("(18.0)", rate=True) == "18.0" and clean_number("(135.00)") == "-135.00"


def test_table_descriptions_above_the_amounts_need_text_before():
    from invoice_engine.table import TableLayoutError
    text = (_HDR + "\n"
            "Shoes          RED TAPE Sneakers\n"
            "                                           1   785.00    -135.00    619.05   30.95    650.00\n"
            "Socks          Cotton\n"
            "                                           1    40.00     -40.00      0.00    0.00      0.00\n" + _END)
    with pytest.raises(TableLayoutError, match="text: before"):
        _table(text)
    rows = _table(text, text="before")
    assert rows[0]["name"] == "RED TAPE Sneakers" and rows[1]["name"] == "Cotton"


def test_table_serial_number_one_space_from_the_description():
    text = ("#  Description                      Qty    Rate     Total\n\n"
            "1 Pigeon Induction Cooktop            1   1499.00   1499.00\n"
            "2 Shipping Charges                    1     40.00     40.00\n"
            "                          Total       2             1539.00\n")
    rows = _table(text, start=r"(?m)^#\s+Description", columns={"pos": "^#", "name": "Description",
                  "qty": "Qty", "price_unit": "Rate", "price_total": "Total"},
                  types={"pos": "int", "qty": "int", "price_unit": "float", "price_total": "float"})
    assert [(r["pos"], r["name"]) for r in rows] == [(1, "Pigeon Induction Cooktop"), (2, "Shipping Charges")]



# --- Indian invoice variations the AI must cover --------------------------------------------------

def test_gst_coverage_of_an_ai_draft():
    from invoice_engine.ai_author import _gst_coverage
    table = {"parser": "table", "start": "x", "end": "y",
             "columns": {"name": "Title", "price_subtotal": "Taxable", "igst_amount": "IGST", "price_total": "Total"}}
    draft = {"options": {"currency": "INR"}, "fields": {"amount": "x", "amount_igst": "y", "lines": table}}
    problems = _gst_coverage(draft)
    assert any("amount_cgst, amount_sgst" in p for p in problems)          # totals: one kind only
    assert any("cgst_amount, sgst_amount" in p for p in problems)          # table: one kind only
    table["columns"].update(cgst_amount="CGST", sgst_amount="SGST")
    draft["fields"].update(amount_cgst="a", amount_sgst="b")
    assert [p for p in _gst_coverage(draft)] == ["the sgst_amount heading should also match UTGST (union "
                                                 "territories): for example 'SGST|UTGST'"]
    table["columns"]["sgst_amount"] = "SGST|UTGST"
    assert _gst_coverage(draft) == []
    # a document in another currency has no GST
    assert _gst_coverage({"options": {"currency": "USD"}, "fields": {"amount_igst": "y"}}) == []
    # row patterns for one kind only
    rows = {"options": {"currency": "INR"}, "fields": {"lines": {"parser": "lines", "start": "a", "end": "b",
            "line": "^(?P<name>.+?)\\s+(?P<igst_amount>[\\d.]+)\\s+(?P<price_total>[\\d.]+)$"}}}
    assert any("one kind of GST" in p for p in _gst_coverage(rows))


def test_parts_printed_on_some_documents_only_become_optional():
    from invoice_engine.ai_author import _empty_fields, _repair_mechanical
    draft = {"required_fields": ["issuer", "invoice_number", "amount", "buyer_tax_id", "seller_company_reg"],
             "fields": {"invoice_number": "a", "amount": "b", "buyer_tax_id": "c", "seller_company_reg": "d",
                        "amount_igst": "e", "amount_cgst": "f", "amount_sgst": "g"},
             "optional_fields": ["amount_cgst", "amount_sgst"]}
    notes = _repair_mechanical(draft)
    assert draft["required_fields"] == ["issuer", "invoice_number", "amount"]
    assert set(draft["optional_fields"]) >= {"buyer_tax_id", "seller_company_reg"} and notes
    # on an IGST document the CGST and SGST totals are empty by nature, the IGST total is not
    found = {"invoice_number": "1", "amount": 10, "buyer_tax_id": "x", "seller_company_reg": "y", "amount_igst": 1}
    assert _empty_fields(draft, {"fields": found}) == []
    del found["amount_igst"]
    assert "amount_igst" in _empty_fields(draft, {"fields": found})[0]


def test_stress_losses_are_sent_back_with_a_hint():
    from invoice_engine.ai_author import _stress_problems
    problems = _stress_problems(["stress: short decimals on a.pdf p1: OK -> FAIL; missing: amount",
                                 "regression: something else"])
    assert len(problems) == 1 and "without decimals" in problems[0] and "missing: amount" in problems[0]


def test_gst_draft_reading_igst_only_is_sent_back_and_the_complete_one_kept(tmp_path: Path):
    pdf = need("OD437334801620974100.pdf")
    types, config = _project(tmp_path, drop=["flipkart_seller_tax_invoice.yml"])
    good = _body("flipkart_seller_tax_invoice.yml")
    igst_only = good.replace("      cgst_amount: 'CGST'\n", "").replace("      sgst_amount: 'SGST|UTGST'\n", "")
    assert igst_only != good
    fake = FakeAI({"document_type": "invoice", "is_new_type": False, "issuer": "Flipkart",
                   "variant_of": None, "template_name": "flipkart_seller_ai"}, [igst_only, good, good, good])
    result = _ai_pipeline(tmp_path, types, config, fake).process(pdf)
    ai = result["ai"][0]
    assert (ai["outcome"], ai["attempts"]) == ("admitted", 2)
    correction = next(p for p in fake.prompts if "rejected" in p)
    assert "cgst_amount, sgst_amount" in correction
    assert [d["status"] for d in result["documents"] if d["template"] == "flipkart_seller_ai.yml"] == ["OK", "OK"]


# --- a rounding rule a first sample cannot settle --------------------------------------------------

def _rounding_check(amount, rows, mode):
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    from invoice_engine.validate import check_document
    base = {"issuer": "x", "document_type": "t", "invoice_number": "1", "date": "2026-01-01", "seller_name": "x"}
    template["fields"] = {k: "x" for k in list(base) + ["amount"]}
    template.update(total_rounding=1, total_rounding_mode=mode)
    issues, evidence = check_document(dict(base, amount=amount),
                                      [{"name": str(i), "price_total": v} for i, v in enumerate(rows)], [], template, REGISTRY)
    return any("sum of line price_total vs amount" in m for _, m in issues), evidence.get("rounding_fits")


def test_rounding_rule_list_passes_with_any_and_reports_which_fit():
    assert _rounding_check(1379.0, [1340.0, 38.6], ["up", "nearest"]) == (False, ["up", "nearest"])
    assert _rounding_check(519.0, [500.0, 18.15], ["up", "nearest"]) == (False, ["up"])
    assert _rounding_check(518.0, [500.0, 18.15], ["up", "nearest"]) == (False, ["nearest"])
    assert _rounding_check(520.0, [500.0, 18.15], ["up", "nearest"])[0] is True          # fits none: error
    assert _rounding_check(519.0, [500.0, 18.15], "up") == (False, None)                 # one rule: no list
    assert _rounding_check(518.5, [500.0, 18.15], ["up", "nearest"]) == (True, None)     # not a whole total


def test_rounding_rule_list_lint():
    from invoice_engine.lint import lint_template
    template = dict(yaml.safe_load(CESS_TEMPLATE), total_rounding=1)
    for mode, ok in (("up", True), (["up", "nearest"], True), (["up", "sideways"], False),
                     ([], False), (["up", "up"], False)):
        errors, _ = lint_template(dict(template, total_rounding_mode=mode), REGISTRY)
        assert (not any("total_rounding_mode" in e for e in errors)) == ok, mode


def test_rounding_hint_names_every_rule_that_fits():
    from invoice_engine.ai_author import _rounding_hint
    def hint(calculated, printed):
        doc = {"issues": [{"message": f"sum of line price_total vs amount: calculated {calculated}, printed {printed}"}]}
        return (_rounding_hint(doc) or [""])[0]
    assert "total_rounding_mode: [up, nearest]" in hint("1378.60", "1379.00")
    assert "total_rounding_mode: up" in hint("518.15", "519.00") and "[" not in hint("518.15", "519.00")
    assert "total_rounding_mode: [up, nearest]" in hint("1378.50", "1379.00")           # half rounds up
    assert "total_rounding_mode: down" in hint("1378.50", "1378.00")
    assert hint("518.15", "520.00") == ""


def test_a_later_document_settles_the_rounding_rule(tmp_path: Path):
    first, second = need("irctc.pdf"), need("2147438163.pdf")              # 1378.60 -> 1379, 518.15 -> 519
    from invoice_engine.pipeline import Pipeline
    from invoice_engine.samples import SampleStore
    types, config = _project(tmp_path)
    path = types / "invoice" / "templates" / "indian_railways_train_ticket.yml"
    path.write_text(path.read_text(encoding="utf-8").replace("total_rounding_mode: up", "total_rounding_mode: [up, nearest]"),
                    encoding="utf-8")
    store = SampleStore(tmp_path / "samples")
    pipe = Pipeline(load_catalog(types, config), store, write_layouts=False)
    result = pipe.process(first)
    assert result["documents"][0]["status"] == "OK" and "[up, nearest]" in path.read_text(encoding="utf-8")
    result = pipe.process(second)
    assert result["documents"][0]["status"] == "OK"
    assert any("settled to up" in n for n in result.get("notes", []))
    text = path.read_text(encoding="utf-8")
    assert "total_rounding_mode: up   # settled" in text and "[up, nearest]" not in text.split("settled")[0]
    # the next run checks the changed template on both stored tickets and reads them as before
    again = Pipeline(load_catalog(types, config), store, write_layouts=False)
    assert [again.process(p)["documents"][0]["status"] for p in (first, second)] == ["OK", "OK"]
    assert again.state["indian_railways_train_ticket.yml"] == "admitted"


# --- text fields no arithmetic can prove ----------------------------------------------------------

def test_gstin_check_character():
    from invoice_engine.textchecks import valid_gstin
    for real in ("27AALCR3173P1ZN", "29AACCF0683K1ZD", "24AAHFT7456J1ZR", "07CIVPC2488R2Z7"):
        assert valid_gstin(real)
    assert not valid_gstin("27AALCR3173P1ZM")          # last character misread
    assert not valid_gstin("27AALCR3I73P1ZN")          # 1 read as I


def test_text_checks_catch_the_misreads_seen_so_far():
    from invoice_engine.textchecks import text_issues
    def messages(**fields):
        return [m for _, m in text_issues(fields)]
    # a seller name that ran into the invoice number; a buyer address with a PAN at its end
    assert "Invoice Number #" in messages(seller_name="ASIAN RETAIL VENTURES , Invoice Number # FAGR6O2400141858")[0]
    assert "PAN: A" in messages(buyer_address="Saijpur Bogha. Ahmedabad 382345 Gujarat PAN: ABAFA8599B")[0]
    assert messages(seller_name="SOLD BY: X")                                      # labels in capitals too
    assert messages(seller_name="Japan: Trading Co", buyer_address="Pancham Society, Kalpana: road") == []
    # formats
    assert "not a valid GSTIN" in messages(seller_tax_id="27AALCR3173P1ZM")[0]
    assert "not a PAN" in messages(seller_pan="AALCR317P")[0]
    assert "CIN" in messages(seller_company_reg="U72900DL2011PLC22561")[0]
    assert "FSSAI" in messages(seller_fssai="1222302600044")[0]
    # agreement between GSTIN, PAN and state
    assert "differs from the PAN inside" in messages(seller_tax_id="27AALCR3173P1ZN", seller_pan="AALCR3137P")[0]
    assert "state" in messages(buyer_tax_id="29AACCU5941P1ZL", buyer_state_code="27")[0]
    assert messages(buyer_tax_id="29AACCU5941P1ZL", buyer_state="Bangalore, Karnataka") == []
    assert messages(seller_tax_id="24AABCF8078M2Z6", seller_state="IN-GJ") == []


def test_text_checks_leave_placeholders_and_other_countries_alone():
    from invoice_engine.textchecks import text_issues
    assert text_issues({"seller_tax_id": "UNREGISTERED", "buyer_tax_id": "NA", "seller_pan": "N.A."}) == []
    assert text_issues({"seller_tax_id": "DE232446240", "seller_company_reg": "HRB 13302"}) == []
    assert text_issues({"seller_tax_id": "NL820471616B01", "seller_company_reg": "32147382"}) == []
    # a partner's personal PAN printed on a firm's invoice is not the firm's PAN: not compared
    assert text_issues({"seller_tax_id": "24AAHFT7456J1ZR", "seller_pan": "AJEPA2477N"}) == []
    # a state the table does not know is not compared
    assert text_issues({"seller_tax_id": "07CIVPC2488R2Z7", "seller_state": "Delhi/DL"}) == []


def test_a_misread_text_field_makes_the_document_check():
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    from invoice_engine.validate import check_document, status_of
    base = {"issuer": "x", "document_type": "t", "invoice_number": "1", "date": "2026-01-01",
            "seller_name": "ASIAN RETAIL VENTURES , Invoice Number # FAGR6O2400141858", "amount": 10.0}
    template["fields"] = {k: "x" for k in base}
    issues, _ = check_document(base, [{"name": "a", "price_total": 10.0}], [], template, REGISTRY)
    assert status_of(issues) == "CHECK" and any("seller_name contains the label" in m for _, m in issues)


# --- both GST kinds: IGST (other state) and CGST + SGST (same state) -------------------------------

_GST_CASES = [   # template, sample pdf, pages, the kind the converted document must show
    ("shiprocket_charges_invoice.yml", "SRC27DL000011945.pdf", [1], "cgst"),
    ("shiprocket_credit_note.yml", "CNSRFW27HR005195.pdf", [1], "cgst"),
    ("shiprocket_services_invoice.yml", "SRV27HR000090197.pdf", [1], "cgst"),
    ("ishwam_sales_invoice.yml", "SALES_1_ISHWAM.pdf", [1], "cgst"),
    ("mtalkz_tax_invoice.yml", "TI26-2701230.pdf", [1, 2], "cgst"),
    ("flipkart_platform_fee_invoice.yml", "OD437334801620974100.pdf", [4], "cgst"),
    ("flipkart_bill_of_supply.yml", "OD437334801620974100.pdf", [6, 7], "igst"),
    ("flipkart_seller_own_invoice.yml", "OD437298764051095100.pdf", [4], "cgst_line"),
    ("meesho_tax_invoice.yml", "ms_invoice_326255343871054720_1.pdf", [1], "cgst_type"),
    ("meesho_credit_note.yml", "ms_invoice_326255343871054720_1.pdf", [2], "cgst_type"),
    ("zomato_platform_fee_invoice.yml", "Invoice_8626225731.pdf", [2], "igst"),
    ("zomato_restaurant_invoice.yml", "Invoice_8626225731.pdf", [1], "igst"),
    ("zomato_restaurant_invoice.yml", "Invoice_8261790801.pdf", [1], "igst"),
    ("blinkit_goods_invoice.yml", "ForwardInvoice_ORD85101206944.pdf", [1], "igst"),
    ("blinkit_fee_invoice.yml", "ForwardInvoice_ORD85101206944.pdf", [2], "igst"),
    ("flipkart_grocery_invoice.yml", "OD435477552035539200.pdf", [4, 5], "igst"),
    ("vodafone_idea_postpaid_bill.yml", "VIL_bill_7573001478_2026-08-22_statement.pdf", [1, 2, 3, 4, 5], "igst"),
]


def _gst_variant(template_name, pdf_name, pages):
    import gst_variants
    from invoice_engine.document import attempt
    from invoice_engine.segment import canonical, page_texts
    pdf = need(pdf_name)
    catalog = load_catalog(ROOT / "document_types", ROOT / "config")
    template = next(t for t in catalog.templates if t["template_name"] == template_name)
    texts = page_texts(pdf)
    original = [canonical(texts[p - 1]) for p in pages]
    converted = [gst_variants.CONVERT[template_name](t) for t in original]
    assert converted != original, "the sample was not converted"
    return template, converted, pdf


@pytest.mark.parametrize("template_name, pdf_name, pages, kind", _GST_CASES)
def test_template_reads_the_other_gst_kind(template_name, pdf_name, pages, kind):
    from invoice_engine.document import attempt
    template, converted, pdf = _gst_variant(template_name, pdf_name, pages)
    # fields read from word positions (parser: region) come from the PDF itself: party blocks
    # are the same in both kinds
    result = attempt(template, converted, pdf, REGISTRY, pages, positions=True)
    issues = [i["message"] for i in result["issues"] if i["level"] != "INFO"]
    assert result["status"] == "OK", issues
    fields, lines = result["fields"], result["lines"]
    if kind == "cgst":
        assert fields.get("amount_cgst") is not None or any("cgst_amount" in l for l in lines)
        assert fields.get("amount_igst") is None and not any(l.get("igst_amount") for l in lines)
    elif kind == "cgst_line":
        assert all(l.get("cgst_amount") is not None and l.get("sgst_amount") is not None for l in lines)
    elif kind == "cgst_type":          # one tax column naming its kind: CGST on the line, SGST below
        assert all(l.get("tax_type") == "CGST" and l.get("sgst_amount") is not None for l in lines), lines
        assert "SGST" not in " ".join(str(l.get("name")) for l in lines)
    else:
        assert fields.get("amount_igst") is not None and fields.get("amount_cgst") is None


@pytest.mark.parametrize("template_name", ["meesho_tax_invoice.yml", "meesho_credit_note.yml"])
def test_meesho_reads_cgst_and_sgst_printed_on_one_line(template_name):
    import gst_variants
    from invoice_engine.document import attempt
    from invoice_engine.segment import canonical, page_texts
    pdf = need("ms_invoice_326255343871054720_1.pdf")
    page = 1 if template_name == "meesho_tax_invoice.yml" else 2
    template = next(t for t in CATALOG.templates if t["template_name"] == template_name)
    text = gst_variants.meesho_same_line(canonical(page_texts(pdf)[page - 1]))
    result = attempt(template, [text], pdf, REGISTRY, [page], positions=True)
    assert result["status"] == "OK", [i["message"] for i in result["issues"] if i["level"] != "INFO"]
    assert all(l.get("tax_type") == "CGST" and l.get("sgst_amount") for l in result["lines"])
    assert "SGST" not in " ".join(str(l.get("name")) for l in result["lines"])   # not read as item text


@pytest.mark.parametrize("template_name, pdf_name, pages, kind", [c for c in _GST_CASES if c[3] == "cgst"])
def test_a_wrong_cgst_value_in_the_other_kind_fails(template_name, pdf_name, pages, kind):
    """The second kind's values are proven by the same arithmetic: one wrong value is not OK."""
    from invoice_engine.document import attempt
    template, converted, pdf = _gst_variant(template_name, pdf_name, pages)
    original = attempt(template, converted, pdf, REGISTRY, pages, positions=True)
    cgst = original["fields"].get("amount_cgst")
    target = f"{cgst:,.2f}" if cgst else None
    if not target or not any(target in t for t in converted):
        pytest.skip("no document-level CGST total printed")
    broken = [t.replace(target, f"{cgst + 1:,.2f}", 1) for t in converted]
    assert not any(i["level"] == "ERROR" for i in original["issues"])
    assert any(i["level"] == "ERROR" for i in attempt(template, broken, pdf, REGISTRY, pages, positions=True)["issues"])


# --- blocks read from word positions (parser: region) ----------------------------------------------

FIXTURES = Path(__file__).parent / "fixtures"


def _region(pdf_name, heading, **kw):
    from invoice_engine import words
    return words.region([words.page_words(FIXTURES / pdf_name)[0]], heading, **kw)


def test_region_keeps_a_crowded_neighbour_column_out():
    # the Bill To line ends one space before the Ship To column: in the page text the two are one
    # space apart, so a regex that cuts a column at two spaces took "Plot 9, Sector 21, Vashi"
    # (the Ship To address) into the Bill To address; the word positions keep them apart
    assert _region("party_block_tight.pdf", "Bill To", stop="^Phone", skip=1) == [
        "Flat 4, Lake View Apartments, MG Road, Kore", "Pune 411001 Maharashtra"]
    assert _region("party_block_tight.pdf", "Ship To", stop="^Phone") == [
        "Arjun Mehta", "Plot 9, Sector 21, Vashi", "Navi Mumbai 400703 Maharashtra"]


def test_region_leaves_out_the_left_column_and_a_note_at_the_right_edge():
    assert _region("party_block.pdf", "Bill To", stop="^Phone", skip=1) == [
        "Flat 4, Lake View, MG Road,", "Pune 411001 Maharashtra"]                      # no PAN line
    assert _region("party_block.pdf", "Ship To", stop="^Phone", skip=1) == [
        "Flat 4, Lake View, MG Road,", "Pune 411001 Maharashtra"]                      # no "*Keep this invoice"
    assert _region("party_block.pdf", "Bill To", max_lines=1) == ["Riya Sharma"]
    assert _region("party_block.pdf", "Sold By:", same_line=True, max_lines=0) == ["Test Retail Private Limited ,"]
    assert _region("party_block.pdf", "No Such Heading") is None


def test_region_field_needs_the_pdf_and_reads_it_through_a_template():
    from invoice2data.extract.invoice_template import InvoiceTemplate
    from invoice_engine import words
    template = InvoiceTemplate({"issuer": "x", "keywords": ["x"], "fields": {}, "template_name": "t.yml"})
    spec = {"parser": "region", "heading": "Bill To", "skip": 1, "stop": "^Phone"}
    assert words.parse(template, "buyer_address", spec, "") is None                    # text alone: no value
    token = words.current.set({"pdf": FIXTURES / "party_block_tight.pdf", "pages": [1], "rewrite": None})
    try:
        assert words.parse(template, "buyer_address", spec, "") == \
            "Flat 4, Lake View Apartments, MG Road, Kore Pune 411001 Maharashtra"
    finally:
        words.current.reset(token)


def test_region_lint():
    from invoice_engine.lint import lint_template
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    template["fields"] = dict(template["fields"], buyer_address={"parser": "region", "skip": -1})
    errors, _ = lint_template(template, REGISTRY)
    assert any("missing 'heading'" in e for e in errors) and any("skip must be" in e for e in errors)


# --- variant tracker -----------------------------------------------------------------------------

from invoice_engine import variants  # noqa: E402


def _sample(family="f", added="2026-01-01", pages=(1,), lines=(), **fields):
    return {"pdf": added, "name": f"{added}.pdf", "family": family, "template": f"{family}.yml", "status": "OK",
            "pages": list(pages), "fields": fields, "lines": list(lines), "added": added}


@pytest.mark.parametrize("fields, lines, expected", [
    ({"amount_igst": 18.0, "amount_cgst": None}, [], {variants.IGST}),
    ({"amount_cgst": 9.0, "amount_sgst": 9.0}, [], {variants.CGST}),
    ({"amount_igst": 0.0}, [], {variants.ZERO}),
    ({}, [{"cgst_amount": 4.5, "sgst_amount": 4.5}], {variants.CGST}),
    ({}, [{"tax_type": "IGST", "line_tax_amount": 8.7}], {variants.IGST}),
    ({}, [{"tax_type": "CGST", "line_tax_amount": 4.0}], {variants.CGST}),
    ({"amount_igst": 1.0, "amount_cgst": 1.0}, [], {variants.MIXED}),
    ({"amount_tax": 21.0}, [{"line_tax_amount": 21.0}], set()),          # foreign VAT: no GST kind
])
def test_variant_gst_kind_comes_from_the_values_read(fields, lines, expected):
    labels = variants.variant_of(_sample(lines=lines, **fields))
    assert labels & {variants.IGST, variants.CGST, variants.MIXED, variants.ZERO} == expected


def test_variant_items_pages_and_extras():
    lines = [{"name": "Phone cover", "qty": 1, "discount": 30.0},
             {"name": "Shipping and Handling Charges", "qty": 1, "cess_amount": 0}]
    labels = variants.variant_of(_sample(pages=(1, 2), lines=lines, amount_round_off=-0.4))
    assert {variants.ITEMS, variants.PAGES, variants.DISCOUNT, variants.ROUND_OFF, variants.CHARGE} <= labels
    assert variants.CESS not in labels and variants.ONE_PAGE not in labels
    # a fee invoice whose every row is a charge has no charge row between items
    fees = [{"name": "Shipping charges"}, {"name": "COD charges"}]
    assert variants.CHARGE not in variants.variant_of(_sample(lines=fees))
    assert variants.ONE_ITEM in variants.variant_of(_sample(lines=fees[:1]))


def _family(name):
    return [t for t in CATALOG.templates if Path(t["template_name"]).stem == name]


def test_readable_variants_come_from_the_template_fields():
    flipkart = variants.readable(_family("flipkart_seller_tax_invoice"))
    assert {variants.IGST, variants.CGST, variants.ITEMS, variants.DISCOUNT, variants.CESS, variants.CHARGE} <= flipkart
    vodafone = variants.readable(_family("irctc_e_ticket"))
    assert variants.CGST not in vodafone and variants.IGST not in vodafone
    meesho = variants.readable(_family("meesho_tax_invoice"))           # one tax column naming its kind
    assert variants.IGST in meesho and variants.CGST in meesho
    irctc = variants.readable(_family("irctc_e_ticket"))                 # no GST, no goods table
    assert not irctc & {variants.IGST, variants.CGST, variants.CHARGE}


def test_coverage_lists_what_to_collect_and_what_cannot_be_read(tmp_path: Path):
    from invoice_engine.samples import SampleStore
    store = SampleStore(tmp_path / "samples")
    pdf = tmp_path / "samples" / "pdfs" / "x.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(b"%PDF")
    for family, fields in (("flipkart_seller_tax_invoice", {"amount_igst": 18.0}),
                           ("indian_railways_train_ticket", {"amount_cgst": 1.0, "amount_sgst": 1.0}),
                           ("shopify_bill", {"amount_igst": 3.0})):
        sample = _sample(family, lines=[{"name": "a", "qty": 1}], **fields)
        sample["pdf"] = "x"
        store.data["samples"].append(sample)
    report = {c.family: c for c in variants.coverage(CATALOG, store)}
    flipkart = report["flipkart_seller_tax_invoice"]
    assert flipkart.seen == {variants.IGST: 1, variants.ONE_ITEM: 1, variants.ONE_PAGE: 1}
    assert flipkart.to_collect == [variants.CGST, variants.ITEMS]
    assert variants.PAGES in flipkart.if_available           # several items not seen yet
    assert report["indian_railways_train_ticket"].to_collect[:1] == [variants.IGST]
    shopify = report["shopify_bill"]                          # variants.yml: never CGST+SGST
    assert variants.CGST in shopify.not_expected and not shopify.cannot_read and not shopify.to_collect[:1] == [variants.CGST]
    assert report["amazon_in_tax_invoice"].samples == 0
    written = variants.write_excel(list(report.values()), tmp_path / "variants.xlsx")
    import pandas as pd
    wanted = pd.read_excel(written, sheet_name="To collect")
    assert ((wanted["Layout"] == "flipkart_seller_tax_invoice") & (wanted["Variant"] == variants.CGST)).any()
    assert not ((wanted["Layout"] == "shopify_bill") & (wanted["Variant"] == variants.CGST)).any()


def test_variant_exception_typos_are_reported(tmp_path: Path):
    types, config = _project(tmp_path)
    (types / "invoice" / "variants.yml").write_text(
        "not_expected:\n  shopify_bill:\n    CGST + SGST: typo\n  no_such_layout:\n    IGST: x\n", encoding="utf-8")
    problems = variants.exception_problems(load_catalog(types, config))
    assert any("'CGST + SGST' is not a variant name" in p for p in problems)
    assert any("no_such_layout is not a template" in p for p in problems)
    assert variants.exception_problems(CATALOG) == []


def test_trim_keeps_the_only_sample_of_a_variant():
    old_cgst = _sample(added="2026-01-01", amount_cgst=9.0, amount_sgst=9.0)
    newer = [_sample(added=f"2026-02-{d:02d}", amount_igst=18.0) for d in range(1, 6)]
    kept = variants.keep_on_trim([old_cgst] + newer, 3)
    assert old_cgst in kept and len(kept) == 3
    assert newer[-1] in kept and newer[0] not in kept       # the oldest of the common kind goes
    assert variants.keep_on_trim(newer, 3) == newer[-3:]


def test_store_notes_problems_and_forgets_them_when_read_ok(tmp_path: Path):
    from invoice_engine.samples import SampleStore
    store = SampleStore(tmp_path / "samples")
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4 test")
    doc = {"pages": [1], "template": "t.yml", "status": "FAIL", "fields": {}, "lines": [],
           "issues": [{"level": "WARN", "message": "w"}, {"level": "ERROR", "message": "sum differs"}]}
    store.note_problem(pdf, doc, "t")
    store.note_problem(pdf, dict(doc, status="CHECK"), "t")         # same pages: replaced
    assert [(p["status"], p["reason"]) for p in store.problems_of("t")] == [("CHECK", "sum differs")]
    assert SampleStore(tmp_path / "samples").problems_of("t")       # saved
    store.add(pdf, dict(doc, status="OK"), "t")
    assert store.problems_of("t") == []


def test_pipeline_notes_the_first_document_of_a_new_variant(tmp_path: Path):
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    first = _pipeline(tmp_path)
    first.process(pdf)                                   # admits the template, stores the document
    stored = next(s for s in first.store.data["samples"] if s["family"] == "zomato_restaurant_invoice")
    # pretend the stored document was an IGST one of other pages: the next read is the first CGST+SGST
    stored["fields"] = {k: v for k, v in stored["fields"].items() if not k.startswith(("amount_cgst", "amount_sgst"))}
    stored["fields"]["amount_igst"] = 5.0
    stored["lines"] = [{k: v for k, v in line.items() if not k.startswith(("cgst", "sgst"))} for line in stored["lines"]]
    stored["pages"] = [99]
    first.store.save()
    again = _pipeline(tmp_path)
    result = again.process(pdf)
    notes = [n for n in result.get("notes", []) if n.startswith("first ")]
    assert notes and "CGST+SGST" in notes[0] and "zomato_restaurant_invoice" in notes[0], result.get("notes")
    assert not [n for n in _pipeline(tmp_path).process(pdf).get("notes", []) if n.startswith("first ")]


# --- second opinion on text fields ---------------------------------------------------------------

from invoice_engine import second_opinion as so  # noqa: E402

_PAGE = """Sold By: Bhumika Highstreet India Private Limited     Invoice Number # LIAALHX270001900
GSTIN - 27AAICB4842N1Z8        Bill To                         Ship To
                               Akash Rajput                    Mehta Traders
                               House no.2, Raj Society,        Plot 9, Sector 21,
                               Saijpur Bogha, Ahmedabad 382345 Vashi, Navi Mumbai 400703
Invoice Date: 18-04-2026       Phone: 9876543210
"""


@pytest.mark.parametrize("name, ours, theirs, outcome", [
    ("buyer_address", "House no.2, Raj Society, Saijpur Bogha, Ahmedabad 382345",
     "House no.2, Raj Society,\nSaijpur Bogha, Ahmedabad 382345, India", "agree"),       # joins, country
    ("buyer_name", "Akash Rajput", "AKASH RAJPUT", "agree"),
    ("buyer_address", "House no.2, Raj Society", "House no.2, Raj Society, Saijpur Bogha, Ahmedabad 382345", "differ"),
    ("buyer_address", "House no.2, Raj Society, Saijpur Bogha, Ahmedabad 382345 Phone: 9876543210",
     "House no.2, Raj Society, Saijpur Bogha, Ahmedabad 382345", "differ"),               # label ran in
    ("buyer_address", "Plot 9, Sector 21, Vashi, Navi Mumbai 400703",
     "House no.2, Raj Society, Saijpur Bogha, Ahmedabad 382345", "differ"),               # Ship To read as Bill To
    ("seller_tax_id", "27AAICB4842N1Z8", "27 AAICB 4842N1Z8", "agree"),
    ("invoice_number", "LIAALHX270001900", "LIAALHX270001900", "agree"),
    ("date", "2026-04-18", "18-04-2026", "agree"),
    ("date", "2026-04-18", "2026-04-08", "differ"),
    ("date", "2026-04-18", "next Tuesday", "ai_unreliable"),
    ("buyer_name", "Akash Rajput", "Akash Kumar Rajput Sharma", "ai_unreliable"),      # not on the page
    ("seller_tax_id", "27AAICB4842N1Z8", "29AAICB4842N1Z8", "ai_unreliable"),           # not on the page
    ("buyer_name", "Akash Rajput", None, "ai_missing"),
    ("buyer_name", None, "Akash Rajput", "template_missed"),
    ("buyer_name", None, None, "agree"),
])
def test_second_opinion_compares_values(name, ours, theirs, outcome):
    assert so.compare(name, ours, theirs, _PAGE).outcome == outcome


def test_second_opinion_value_read_into_another_field_is_not_a_miss():
    # the AI puts the buyer's name under buyer_trade_name: the template has it as buyer_name
    finding = so.compare("buyer_trade_name", None, "Akash Rajput", _PAGE, ["Akash Rajput"])
    assert finding.outcome == "ai_unreliable" and finding.level == "INFO"


def test_second_opinion_asks_text_and_date_fields_only():
    template = next(t for t in CATALOG.templates if t["template_name"] == "flipkart_seller_tax_invoice.yml")
    names = so.fields_to_check(template, CATALOG.registry_for(template))
    assert {"invoice_number", "date", "seller_name", "seller_tax_id", "buyer_name", "buyer_address"} <= set(names)
    assert not {"amount", "amount_igst", "document_type", "issuer", "currency", "lines"} & set(names)
    question = so.prompt("page text", names)
    assert "LIAALHX" not in question and "page text" in question      # the template's values are not sent


class OpinionAI:
    """Answers the second-opinion prompt with the values a plain read gives, per page, changed
    where asked; counts the questions."""

    def __init__(self, truth: dict[int, dict], change: dict[int, dict] | None = None, fail: bool = False):
        self.truth, self.change, self.fail, self.calls = truth, change or {}, fail, 0

    def complete(self, system: str, prompt: str) -> str:
        from invoice_engine.ai_provider import AIError
        self.calls += 1
        if self.fail:
            raise AIError("service unavailable")
        page = int(re.search(r"=== page (\d+) ===", prompt).group(1))
        names = re.findall(r"(?m)^- (\w+): ", prompt.split("DOCUMENT TEXT")[0])
        values = {n: self.truth[page].get(n) for n in names}
        values = {n: (v.isoformat() if hasattr(v, "isoformat") else v) for n, v in values.items()}
        values.update(self.change.get(page, {}))
        return "```json\n" + json.dumps(values) + "\n```"


def _opinion_pipeline(tmp_path: Path, ai, limit: int = 3):
    from invoice_engine.ai_author import TemplateAuthor
    from invoice_engine.ai_provider import AISettings
    from invoice_engine.pipeline import Pipeline
    from invoice_engine.samples import SampleStore
    store = SampleStore(tmp_path / "samples")
    author = TemplateAuthor(ai, AISettings(enabled=True, model="fake", second_opinion=limit), ROOT / "document_types",
                            ROOT / "config", store)
    return Pipeline(CATALOG, store, write_layouts=False, author=author)


def _zomato_truth() -> tuple[Path, dict]:
    pdf = ROOT / "Invoice_8626225731.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    docs = process_pdf(pdf, CATALOG)["documents"]
    return pdf, {d["pages"][0]: d["fields"] for d in docs}


def test_second_opinion_that_agrees_keeps_documents_ok(tmp_path: Path):
    pdf, truth = _zomato_truth()
    ai = OpinionAI(truth)
    pipe = _opinion_pipeline(tmp_path, ai)
    result = pipe.process(pdf)
    assert [d["status"] for d in result["documents"]] == ["OK", "OK"] and ai.calls == 2
    assert len(pipe.store.data["samples"]) == 2
    assert any("documents of zomato_restaurant_invoice.yml agreed" in i["message"] for i in result["documents"][0]["issues"])
    _opinion_pipeline(tmp_path, ai).process(pdf)                  # the same file again: not asked again
    assert ai.calls == 2


def test_second_opinion_difference_makes_check_and_keeps_it_out_of_the_samples(tmp_path: Path):
    pdf, truth = _zomato_truth()
    swapped = {1: {"seller_name": truth[1]["buyer_name"]}}        # a value from the page, of another party
    ai = OpinionAI(truth, swapped)
    pipe = _opinion_pipeline(tmp_path, ai)
    result = pipe.process(pdf)
    restaurant, fee = result["documents"]
    assert restaurant["status"] == "CHECK" and fee["status"] == "OK"
    assert any(i["level"] == "WARN" and "seller_name differs" in i["message"] for i in restaurant["issues"])
    assert any("--trust zomato_restaurant_invoice.yml" in i["message"] for i in restaurant["issues"])
    assert {s["family"] for s in pipe.store.data["samples"]} == {"zomato_platform_fee_invoice"}
    assert pipe.store.problems_of("zomato_restaurant_invoice")      # shown in the variant report
    again = _opinion_pipeline(tmp_path, ai).process(pdf)            # read again: same result, no new question
    assert again["documents"][0]["status"] == "CHECK" and ai.calls == 2
    # a person checked it: the template is right
    from invoice_engine.admission import template_hash
    template = next(t for t in CATALOG.templates if t["template_name"] == "zomato_restaurant_invoice.yml")
    so.trust(pipe.store, "zomato_restaurant_invoice.yml", template_hash(CATALOG, template))
    trusted = _opinion_pipeline(tmp_path, ai)
    assert [d["status"] for d in trusted.process(pdf)["documents"]] == ["OK", "OK"] and ai.calls == 2
    assert "zomato_restaurant_invoice" in {s["family"] for s in trusted.store.data["samples"]}


def test_second_opinion_unavailable_leaves_documents_ok_and_asks_again_later(tmp_path: Path):
    pdf, truth = _zomato_truth()
    pipe = _opinion_pipeline(tmp_path, OpinionAI(truth, fail=True))
    result = pipe.process(pdf)
    assert [d["status"] for d in result["documents"]] == ["OK", "OK"]
    assert any("second opinion not available" in i["message"] for i in result["documents"][0]["issues"])
    assert not pipe.store.data.get("second_opinions")
    assert len(pipe.store.data["samples"]) == 2                    # stored as before the second opinion
    working = OpinionAI(truth, {1: {"seller_name": truth[1]["buyer_name"]}})
    again = _opinion_pipeline(tmp_path, working).process(pdf)       # the AI answers now: asked again
    assert working.calls == 2 and [d["status"] for d in again["documents"]] == ["CHECK", "OK"]
    store = _opinion_pipeline(tmp_path, working).store             # the doubted document left the store
    assert {s["family"] for s in store.data["samples"]} == {"zomato_platform_fee_invoice"}


def test_second_opinion_stops_after_the_limit_and_restarts_when_the_template_changes(tmp_path: Path):
    from invoice_engine.samples import SampleStore
    store = SampleStore(tmp_path / "samples")
    agree = so.Opinion(findings=[so.Finding("buyer_name", "agree")])
    for n in range(2):
        so.record(store, "t.yml", "hash1", f"doc{n}:1", agree, f"doc{n}")
    assert so.wanted(store, "t.yml", "hash1", 3)
    so.record(store, "t.yml", "hash1", "doc2:1", agree, "doc2")
    assert not so.wanted(store, "t.yml", "hash1", 3)
    assert so.wanted(store, "t.yml", "hash2", 3)                   # the file changed: start again
    assert not so.wanted(store, "t.yml", "hash1", 0)               # switched off
    so.trust(store, "u.yml", "h")
    assert not so.wanted(store, "u.yml", "h", 3) and so.wanted(store, "u.yml", "h2", 3)
    differ = so.Opinion(findings=[so.Finding("buyer_name", "differ", "a", "b")])
    so.record(store, "v.yml", "h", "d:1", differ, "d")
    assert so.agreed_count(store, "v.yml", "h") == 0 and so.wanted(store, "v.yml", "h", 3)


def test_second_opinion_on_stored_samples(tmp_path: Path):
    pdf, truth = _zomato_truth()
    _pipeline(tmp_path).process(pdf)                               # stored without a second opinion
    ai = OpinionAI(truth, {2: {"buyer_address": truth[2]["seller_address"]}})     # on its page, another party
    pipe = _opinion_pipeline(tmp_path, ai)
    report = so.check_samples(CATALOG, pipe.store, pipe.author)
    verdicts = {r["template"]: r["agreed"] for r in report}
    assert verdicts == {"zomato_restaurant_invoice.yml": True, "zomato_platform_fee_invoice.yml": False}
    assert ai.calls == 2 and so.check_samples(CATALOG, pipe.store, pipe.author) == []     # not asked twice


# --- regions: more block shapes ------------------------------------------------------------------

def test_region_reads_a_block_whose_neighbour_column_is_half_a_line_lower():
    # Shiprocket-style: State Code / Place of Supply / GSTIN printed between the buyer's lines
    assert _region("staggered_block.pdf", "^Invoice To:", max_lines=1) == ["MEHTA TRADING COMPANY PRIVATE LIMITED"]
    assert _region("staggered_block.pdf", "^Invoice To:", skip=1) == [
        "Plot 9, Sector 21, Opposite Railway Station Road,", "Vashi, Navi Mumbai", "Maharashtra - 400703"]


def test_region_reads_a_label_value_block_up_to_the_next_label():
    # Blinkit-style: "Address : value" (the colon drawn as a word of its own), lines hanging under
    # the value, "Place of Supply:" one space after the value on the label's line, Pincode: next
    assert _region("label_value_block.pdf", "^Name:", same_line=True, max_lines=1) == ["Riya Sharma"]
    assert _region("label_value_block.pdf", "^Address:", same_line=True) == [
        "Flat 12, Shanti Apartments, Near Old Water Tank Road", "Navrangpura, Ahmedabad, Gujarat, 380009"]
    # the page text has the neighbour one space away: a two-space column cut keeps it in
    from invoice_engine.segment import page_texts
    text = page_texts(FIXTURES / "label_value_block.pdf")[0]
    old = re.search(r"(?s)\nAddress\s*:\s*(.*?)\nPincode", "\n" + text).group(1)
    assert "Place of Supply" in re.sub(r"(?m)(?<=\S)[ \t]{3,}\S[^\n]*$", "", old)


def test_region_join_and_settings_are_checked():
    from invoice_engine import words
    from invoice_engine.lint import lint_template
    token = words.current.set({"pdf": FIXTURES / "staggered_block.pdf", "pages": [1], "rewrite": None})
    try:
        value = words.parse(None, "buyer_address", {"heading": "^Invoice To:", "skip": 1, "join": ", "}, "")
    finally:
        words.current.reset(token)
    assert value == "Plot 9, Sector 21, Opposite Railway Station Road,, Vashi, Navi Mumbai, Maharashtra - 400703"
    template = dict(yaml.safe_load(CESS_TEMPLATE))
    template["fields"] = dict(template["fields"], buyer_address={"parser": "region", "heading": "Bill To",
                                                                 "stops": "^Phone", "join": 1})
    errors, _ = lint_template(template, REGISTRY)
    assert any("unknown setting(s) stops" in e for e in errors) and any("join must be a text" in e for e in errors)



@pytest.mark.parametrize("template_name, pdf_name, pages, kind", [c for c in _GST_CASES if c[3] == "igst"])
def test_a_wrong_igst_value_in_the_other_kind_fails(template_name, pdf_name, pages, kind):
    """An IGST print with one wrong tax value is not OK (the same arithmetic proves the IGST kind)."""
    from invoice_engine.document import attempt
    template, converted, pdf = _gst_variant(template_name, pdf_name, pages)
    original = attempt(template, converted, pdf, REGISTRY, pages, positions=True)
    igst = original["fields"].get("amount_igst")
    assert igst is not None, "no IGST total read"
    if not igst:
        pytest.skip("zero tax (bill of supply): nothing to change")
    target = next((f for f in (f"{igst:,.2f}", f"{igst:.2f}", f"{igst:.3f}") if any(f in t for t in converted)), None)
    if target is None:
        pytest.skip("IGST total printed in another form")
    # change the printed value the template reads as amount_igst (the same number may also be
    # printed in a summary or subtotal row the template does not read)
    tried = 0
    for page, text in enumerate(converted):
        for found in re.finditer(re.escape(target), text):
            broken = list(converted)
            broken[page] = text[:found.start()] + f"{igst + 1:,.2f}" + text[found.end():]
            result = attempt(template, broken, pdf, REGISTRY, pages, positions=True)
            if result["fields"].get("amount_igst") == original["fields"].get("amount_igst"):
                continue
            tried += 1
            assert result["status"] != "OK", [i["message"] for i in result["issues"]]
    assert tried, "no printed value changed amount_igst"
